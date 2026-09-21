"""
PopPilot — génération des rapports réglementaires (BCC).

TROIS RAPPORTS, et ils ne se ressemblent pas :

  - FINA : remplit le gabarit .xls de la BCC depuis le socle. Le gabarit est
    fourni à chaque génération — c'est un document officiel, versionné par la
    BCC, que la plateforme n'a pas à stocker ni à deviner.
  - AML / LBC-FT : le classeur fourni est À LA FOIS la source (brouillards,
    grand livre) et le document à remplir. Rien de ces données n'est conservé :
    la copie de travail est détruite à la fin.
  - SYSTÈME DE PAIEMENT : calcul seulement. Le dépôt ne contient AUCUN
    remplisseur de gabarit pour ce rapport — on produit donc un classeur de
    résultats, clairement nommé comme tel, et on le DIT plutôt que de laisser
    croire que la déclaration officielle est produite.

CE QUI N'EST JAMAIS FAIT ICI : figer un taux. FINA et AML convertissent en CDF
au taux saisi pour la période ; s'il manque, la génération ÉCHOUE avec le
message qui dit quoi saisir. Un taux figé produit un rapport faux d'un facteur
~2268 sans lever la moindre alerte, et ce rapport part à la banque centrale.

LES FEUILLES NON CONCERNÉES (F4a, F4b, F8, F9, F12) ne sont jamais touchées :
elles restent telles que la BCC les a livrées.
"""
from __future__ import annotations

import datetime as dt
import os
import shutil
import tempfile
from dataclasses import dataclass
from typing import Callable

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse
from starlette.background import BackgroundTask

from auth_supabase import ROLES_ACCES_TOTAL, exiger_role, utilisateur_courant
from import_cbs import (ERREURS_DE_FICHIER, _ecrire_sur_disque,
                        _est_marqueur_fastapi, _nom_sain)

routeur = APIRouter(tags=["rapports"])

TABLEUR = (".xls", ".xlsx", ".xlsm")


@dataclass(frozen=True)
class Champ:
    """Un paramètre de génération, décrit pour que le front construise son champ."""
    nom: str
    libelle: str
    type: str                     # "date" | "texte" | "nombre" | "entier"
    obligatoire: bool
    aide: str = ""


@dataclass(frozen=True)
class Fichier:
    nom: str
    libelle: str
    extensions: tuple[str, ...]
    obligatoire: bool
    aide: str = ""


@dataclass(frozen=True)
class Rapport:
    libelle: str
    description: str
    fichiers: tuple[Fichier, ...]
    champs: tuple[Champ, ...]
    #  Extension du document produit : elle suit celle du gabarit, pas un choix.
    #  Le gabarit FINA de la BCC n'existe qu'en .xls (BIFF), qu'openpyxl ne sait
    #  pas écrire — d'où xlutils/xlwt côté moteur.
    extension_sortie: str
    #  Le gabarit officiel est-il rempli, ou produit-on un classeur de résultats ?
    remplit_gabarit: bool
    aide: str


RAPPORTS: dict[str, Rapport] = {
    "fina": Rapport(
        libelle="FINA — états périodiques BCC",
        description="Remplit le gabarit MFII de la BCC depuis le socle, en CDF.",
        fichiers=(
            Fichier("gabarit", "Gabarit BCC (MFII*.xls)", (".xls",), True,
                    "Le gabarit officiel livré par la BCC. Il est copié, jamais modifié."),
        ),
        champs=(
            Champ("arrete", "Date d'arrêté", "date", True,
                  "Les 9 feuilles sont calculées à cet arrêté."),
            Champ("nb_employes", "Effectif total (F2)", "entier", False,
                  "Rapport RH mensuel. Sans lui, la case reste vide."),
            Champ("nb_agents_credit", "Dont agents de crédit (F2)", "entier", False,
                  "Effectif RH officiel, pas le roster du fichier OBJECTIF."),
            Champ("f10_commerce", "F10 — part commerce", "nombre", False,
                  "Fraction entre 0 et 1 (0,81 = 81 %). La somme des quatre doit faire 1."),
            Champ("f10_agricole", "F10 — part agricole", "nombre", False, ""),
            Champ("f10_services", "F10 — part services", "nombre", False, ""),
            Champ("f10_autres", "F10 — part autres", "nombre", False, ""),
        ),
        extension_sortie=".xls",
        remplit_gabarit=True,
        aide="Feuilles remplies : F0, F1, F2, F3, F5, F6, F7, F10, F11. "
             "JAMAIS touchées : F4a, F4b, F8, F9, F12. DEUX PRÉALABLES : la balance "
             "CDF de l'arrêté doit être importée (le FINA se déclare en CDF, et la "
             "balance USD ne la remplace pas), et un taux USD→CDF doit être saisi "
             "pour la période (page Configuration).",
    ),
    "aml": Rapport(
        libelle="AML / LBC-FT — déclaration anti-blanchiment",
        description="Remplit le rapport LBC-FT : opérations espèces, transferts, "
                    "portefeuille client, localisation.",
        fichiers=(
            Fichier("source", "Classeur LBC-FT (gabarit + brouillards + grand livre)",
                    TABLEUR, True,
                    "Ce classeur est à la fois la SOURCE des opérations et le document "
                    "à remplir. La copie de travail est détruite après génération."),
            Fichier("inventaire", "Inventaire dépôt de la période", (".csv",) + TABLEUR,
                    False,
                    "Alimente le portefeuille client et la localisation. Sans lui, "
                    "ces sections restent vides."),
        ),
        champs=(
            Champ("periode_debut", "Début de période", "date", True, ""),
            Champ("periode_fin", "Fin de période", "date", True,
                  "Le taux CDF est lu à CETTE date (§42)."),
            Champ("encours_credit_usd", "Encours crédit (USD)", "nombre", False,
                  "Fourni par le CDG."),
            Champ("nb_credits", "Nombre de crédits", "entier", False, ""),
            Champ("credit_conso_cdf", "Crédits à la consommation (CDF)", "nombre", False,
                  "Staff + avance salaire."),
            Champ("nb_credits_conso", "Nombre de crédits conso", "entier", False, ""),
            Champ("ligne_groupe", "Ligne « groupes » dans le gabarit", "entier", False,
                  "Numéro de ligne du statut juridique 4, à relever une fois sur le "
                  "gabarit. Sans lui, les groupes sont comptés mais pas écrits."),
        ),
        extension_sortie=".xlsx",
        remplit_gabarit=True,
        aide="Exige un taux USD→CDF saisi à la date de fin de période. Le grand "
             "livre est toujours en USD : il est converti au taux daté.",
    ),
    "systeme_paiement": Rapport(
        libelle="Système de paiement — comptes actifs et transactions",
        description="Comptes actifs/dormants et types de transactions, par devise SÉPARÉE.",
        fichiers=(
            Fichier("dormant", "Fichier des comptes dormants", TABLEUR, True,
                    "Un compte est actif si sa dernière opération date de moins de "
                    "6 mois avant l'arrêté."),
            Fichier("inventaire", "Inventaire dépôt", (".csv",) + TABLEUR, True,
                    "Source des versements (cash in) et retraits (cash out)."),
        ),
        champs=(
            Champ("arrete", "Date d'arrêté", "date", True,
                  "Détermine la fenêtre de 6 mois pour l'activité des comptes."),
        ),
        extension_sortie=".xlsx",
        remplit_gabarit=False,
        aide="ATTENTION : ce rapport produit un classeur de RÉSULTATS, il ne remplit "
             "pas le gabarit officiel de la BCC — le dépôt ne contient pas de "
             "remplisseur pour ce rapport. Les montants sont ventilés par devise "
             "SANS conversion (contrairement à l'AML, qui convertit).",
    ),
}


def _date(valeur: str, champ: str) -> dt.date:
    try:
        return dt.date.fromisoformat(str(valeur))
    except (ValueError, TypeError):
        raise HTTPException(400, f"{champ} : date invalide ({valeur!r}), format AAAA-MM-JJ.")


def _nombre(valeur, champ: str) -> float | None:
    if valeur is None or (isinstance(valeur, str) and not valeur.strip()):
        return None
    try:
        return float(str(valeur).replace(",", "."))
    except ValueError:
        raise HTTPException(400, f"{champ} : nombre invalide ({valeur!r}).")


def _entier(valeur, champ: str) -> int | None:
    n = _nombre(valeur, champ)
    return None if n is None else int(n)


def _exiger_champs(cle: str, spec: Rapport, valeurs: dict) -> None:
    manquants = [c.libelle for c in spec.champs
                 if c.obligatoire and not str(valeurs.get(c.nom) or "").strip()]
    if manquants:
        raise HTTPException(
            422, f"Paramètre(s) obligatoire(s) manquant(s) pour « {cle} » : "
                 f"{', '.join(manquants)}.")


@routeur.get("/rapports")
def catalogue(user: dict = Depends(utilisateur_courant)):
    """Ce que la plateforme sait produire, et ce qu'il faut lui fournir.

    Le formulaire du front est CONSTRUIT à partir de cette réponse : une seule
    liste à tenir, et elle vit du côté qui sait de quoi les moteurs ont besoin.
    """
    exiger_role(user, ROLES_ACCES_TOTAL)
    return {
        "rapports": [
            {
                "cle": cle,
                "libelle": r.libelle,
                "description": r.description,
                "aide": r.aide,
                "remplit_gabarit": r.remplit_gabarit,
                "extension_sortie": r.extension_sortie,
                "fichiers": [vars(f) for f in r.fichiers],
                "champs": [vars(c) for c in r.champs],
            }
            for cle, r in RAPPORTS.items()
        ]
    }


def _telecharger(chemin: str, nom: str, dossier: str) -> FileResponse:
    """Renvoie le document produit, puis efface le dossier de travail.

    Le nettoyage passe par `BackgroundTask` : supprimer avant l'envoi couperait
    le téléchargement, et ne pas supprimer laisserait sur le disque du serveur
    des brouillards et un grand livre — des données clients.
    """
    return FileResponse(
        chemin,
        filename=nom,
        media_type="application/octet-stream",
        background=BackgroundTask(shutil.rmtree, dossier, ignore_errors=True),
    )


def _recevoir(dossier: str, spec: Rapport, envoyes: dict[str, UploadFile | None]) -> dict:
    """Écrit les fichiers reçus sur disque, en appliquant les mêmes garde-fous
    que l'import (nom assaini, plafond de taille, écriture par morceaux)."""
    chemins: dict[str, str] = {}
    for f in spec.fichiers:
        envoi = envoyes.get(f.nom)
        if envoi is None or not getattr(envoi, "filename", ""):
            if f.obligatoire:
                raise HTTPException(422, f"Fichier obligatoire manquant : {f.libelle}.")
            continue
        nom = _nom_sain(envoi.filename, f.extensions)
        chemin = os.path.join(dossier, f"{f.nom}_{nom}")
        _ecrire_sur_disque(envoi, chemin)
        chemins[f.nom] = chemin
    return chemins


def _executer(cle: str, generer: Callable[[str], dict], dossier: str,
              nom_sortie: str) -> FileResponse:
    """Lance la génération et traduit ses échecs en réponses lisibles."""
    sortie = os.path.join(dossier, nom_sortie)
    try:
        generer(sortie)
    except HTTPException:
        shutil.rmtree(dossier, ignore_errors=True)
        raise
    except ERREURS_DE_FICHIER as e:
        shutil.rmtree(dossier, ignore_errors=True)
        #  Le gestionnaire ValueError global traduirait ceci en 404 « donnée
        #  absente » : ici la donnée n'est pas absente, c'est le fichier fourni
        #  qui ne convient pas, ou une donnée du socle qui manque.
        raise HTTPException(400, f"Génération « {cle} » impossible : {e}")
    except Exception as e:                                   # noqa: BLE001
        shutil.rmtree(dossier, ignore_errors=True)
        raise HTTPException(500, f"Génération « {cle} » interrompue : {type(e).__name__}: {e}")

    if not os.path.exists(sortie):
        shutil.rmtree(dossier, ignore_errors=True)
        raise HTTPException(500, f"Génération « {cle} » : aucun document produit.")
    return _telecharger(sortie, nom_sortie, dossier)


@routeur.post("/rapports/{cle}")
def generer(cle: str,
            gabarit: UploadFile | None = File(None),
            source: UploadFile | None = File(None),
            inventaire: UploadFile | None = File(None),
            dormant: UploadFile | None = File(None),
            arrete: str | None = Form(None),
            periode_debut: str | None = Form(None),
            periode_fin: str | None = Form(None),
            nb_employes: str | None = Form(None),
            nb_agents_credit: str | None = Form(None),
            f10_commerce: str | None = Form(None),
            f10_agricole: str | None = Form(None),
            f10_services: str | None = Form(None),
            f10_autres: str | None = Form(None),
            encours_credit_usd: str | None = Form(None),
            nb_credits: str | None = Form(None),
            credit_conso_cdf: str | None = Form(None),
            nb_credits_conso: str | None = Form(None),
            ligne_groupe: str | None = Form(None),
            user: dict = Depends(utilisateur_courant)):
    """Produit un rapport réglementaire et le renvoie en téléchargement.

    Rien n'est écrit dans le socle : une génération LIT le socle et les fichiers
    fournis. C'est pourquoi elle est ouverte aux rôles à accès total, AUDIT
    compris — recalculer une déclaration pour la vérifier est le travail d'un
    contrôleur, et cela ne modifie rien.
    """
    exiger_role(user, ROLES_ACCES_TOTAL)

    spec = RAPPORTS.get(cle)
    if spec is None:
        raise HTTPException(
            404, f"Rapport inconnu : {cle!r}. Rapports : {', '.join(RAPPORTS)}.")

    #  Appelé par HTTP, FastAPI résout chaque `Form(None)` en `None`. Appelé
    #  DIRECTEMENT — ce que font les suites de tests —, le défaut reste l'objet
    #  `FieldInfo`, qui est TRUTHY : le contrôle des champs obligatoires le prend
    #  pour une saisie et laisse passer, puis la conversion échoue sur un « date
    #  invalide (Form(None)) » au lieu de dire ce qui manque vraiment.
    def _reel(valeur):
        return None if _est_marqueur_fastapi(valeur) else valeur

    gabarit, source = _reel(gabarit), _reel(source)
    inventaire, dormant = _reel(inventaire), _reel(dormant)

    valeurs = {cle_champ: _reel(v) for cle_champ, v in {
        "arrete": arrete, "periode_debut": periode_debut, "periode_fin": periode_fin,
        "nb_employes": nb_employes, "nb_agents_credit": nb_agents_credit,
        "f10_commerce": f10_commerce, "f10_agricole": f10_agricole,
        "f10_services": f10_services, "f10_autres": f10_autres,
        "encours_credit_usd": encours_credit_usd, "nb_credits": nb_credits,
        "credit_conso_cdf": credit_conso_cdf, "nb_credits_conso": nb_credits_conso,
        "ligne_groupe": ligne_groupe,
    }.items()}
    _exiger_champs(cle, spec, valeurs)

    dossier = tempfile.mkdtemp(prefix="poppilot_rapport_")
    try:
        fichiers = _recevoir(dossier, spec, {
            "gabarit": gabarit, "source": source,
            "inventaire": inventaire, "dormant": dormant,
        })
    except HTTPException:
        shutil.rmtree(dossier, ignore_errors=True)
        raise

    if cle == "fina":
        return _fina(dossier, fichiers, valeurs)
    if cle == "aml":
        return _aml(dossier, fichiers, valeurs)
    return _systeme_paiement(dossier, fichiers, valeurs)


# ─────────────────────────────────────────────────────────────────────────────
def _fina(dossier: str, fichiers: dict, v: dict) -> FileResponse:
    from engine.fina_ecriture import ecrire_fina

    d = _date(v["arrete"], "Date d'arrêté")

    rh = {}
    if v["nb_employes"]:
        rh["nb_employes"] = _entier(v["nb_employes"], "Effectif total")
    if v["nb_agents_credit"]:
        rh["nb_agents_credit"] = _entier(v["nb_agents_credit"], "Agents de crédit")

    parts = {cle: _nombre(v[f"f10_{cle}"], f"F10 {cle}")
             for cle in ("commerce", "agricole", "services", "autres")}
    ventilation = {c: p for c, p in parts.items() if p is not None}
    if ventilation:
        somme = sum(ventilation.values())
        #  Le moteur contrôle déjà la somme et alerte ; on refuse ici en amont,
        #  parce qu'une ventilation qui ne fait pas 100 % répartit un total réel
        #  sur des parts qui n'en couvrent qu'une fraction — le total de F10 ne
        #  correspondrait plus au crédit déclaré ailleurs dans le même rapport.
        if abs(somme - 1.0) > 0.005:
            raise HTTPException(
                422, f"Ventilation sectorielle F10 : la somme des parts vaut {somme:.4f} "
                     f"au lieu de 1. Saisir des fractions (0,81 = 81 %) qui totalisent 1, "
                     f"ou n'en saisir aucune.")

    nom = f"FINA_{d.isoformat()}.xls"
    return _executer(
        "fina",
        lambda sortie: ecrire_fina(d, fichiers["gabarit"], sortie,
                                   rh=rh or None,
                                   ventilation_f10=ventilation or None),
        dossier, nom)


def _aml(dossier: str, fichiers: dict, v: dict) -> FileResponse:
    from engine.aml_ecriture import ecrire_aml

    debut = _date(v["periode_debut"], "Début de période")
    fin = _date(v["periode_fin"], "Fin de période")
    if fin < debut:
        raise HTTPException(422, "La fin de période est antérieure à son début.")

    nom = f"AML_LBC-FT_{debut.isoformat()}_{fin.isoformat()}.xlsx"
    return _executer(
        "aml",
        lambda sortie: ecrire_aml(
            fichiers["source"], sortie, debut, fin,
            path_inventaire=fichiers.get("inventaire"),
            encours_credit_usd=_nombre(v["encours_credit_usd"], "Encours crédit"),
            nb_credits=_entier(v["nb_credits"], "Nombre de crédits"),
            credit_conso_cdf=_nombre(v["credit_conso_cdf"], "Crédits conso"),
            nb_credits_conso=_entier(v["nb_credits_conso"], "Nombre crédits conso"),
            ligne_groupe=_entier(v["ligne_groupe"], "Ligne groupes"),
        ),
        dossier, nom)


def _systeme_paiement(dossier: str, fichiers: dict, v: dict) -> FileResponse:
    """Classeur de RÉSULTATS — le gabarit officiel n'est pas rempli (cf. RAPPORTS)."""
    from engine.systeme_paiement import comptes_actifs, transactions_inventaire
    from export_excel import (BLEU, BLANC_GRAS, ENTIER, MONTANT, _entete,
                              _largeurs, _ligne)
    from openpyxl import Workbook

    d = _date(v["arrete"], "Date d'arrêté")
    nom = f"Systeme_paiement_{d.isoformat()}.xlsx"

    def produire(sortie: str) -> dict:
        comptes = comptes_actifs(fichiers["dormant"], d)
        #  PAS de taux : ce rapport ventile par devise SANS convertir. Passer un
        #  taux ici mélangerait les deux doctrines (celle de l'AML et la sienne).
        transactions = transactions_inventaire(fichiers["inventaire"], taux_cdf=None)

        wb = Workbook()
        ws = wb.active
        ws.title = "Comptes"
        ws["A1"] = "Système de paiement — comptes actifs et dormants"
        ws["A1"].font = BLANC_GRAS
        ws["A1"].fill = BLEU
        ws["A2"] = (f"Arrêté du {d.strftime('%d/%m/%Y')} — un compte est ACTIF si sa "
                    f"dernière opération date de moins de 6 mois avant cet arrêté.")
        n = _entete(ws, 4, ["Catégorie", "Nombre"])
        for libelle, cle in (("Total comptes", "total"), ("Actifs", "actifs"),
                             ("Dormants", "dormants")):
            n = _ligne(ws, n, [libelle, comptes.get(cle)], [None, ENTIER])
        for cle, valeur in comptes.items():
            if cle not in ("total", "actifs", "dormants"):
                n = _ligne(ws, n, [cle, valeur], [None, ENTIER])
        _largeurs(ws, [34, 16])

        ws2 = wb.create_sheet("Transactions")
        ws2["A1"] = "Types de transactions, PAR DEVISE — aucune conversion"
        ws2["A1"].font = BLANC_GRAS
        ws2["A1"].fill = BLEU
        ws2["A2"] = ("Contrairement à l'AML, ce rapport ne convertit pas : les montants "
                     "CDF restent en CDF, les USD en USD. Aucun total toutes devises.")
        m = _entete(ws2, 4, ["Type", "Devise", "Nombre", "Montant (devise d'origine)"])
        for type_op in ("versement", "retrait"):
            bloc = transactions.get(type_op) or {}
            for devise, detail in sorted(bloc.items()):
                if isinstance(detail, dict):
                    m = _ligne(ws2, m, [type_op, devise, detail.get("nombre"),
                                        detail.get("montant")],
                               [None, None, ENTIER, MONTANT])
                else:
                    m = _ligne(ws2, m, [type_op, devise, None, detail],
                               [None, None, None, MONTANT])
        _largeurs(ws2, [18, 12, 14, 28])

        wb.save(sortie)
        return {"sortie": sortie}

    return _executer("systeme_paiement", produire, dossier, nom)
