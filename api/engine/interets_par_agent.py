"""
Moteur : intérêts encaissés par agent / superviseur / agence (chantiers 1, 2).
Idée (CDG) : le fichier des remboursements a le n° de dossier + intérêts remboursés,
mais PAS la hiérarchie. L'encours crédit a la hiérarchie (agent, superviseur, agence).
→ jointure sur le numéro de dossier → intérêts ventilés par agent.

Réutilise la logique de rattachement en cascade déjà validée (recouvrement) :
  dossier → encours du mois → sinon orphelin.

NE REMPLACE RIEN : moteur complémentaire, à placer à côté des moteurs existants.
"""
from __future__ import annotations
import datetime as dt
from collections import defaultdict
import openpyxl


def _f(v):
    if v in (None, ""):
        return 0.0
    try:
        return float(str(v).replace(",", ".").replace("\xa0", "").replace(" ", ""))
    except (ValueError, TypeError):
        return 0.0


def lire_remboursements(path, feuille="Worksheet"):
    """Lit le fichier des crédits remboursés (table plate).
    Colonnes : Date, N° échéance, N° client, Noms, N° dossier, Montant déboursé,
    Capital remb., Intérêts remb., Pénalités remb."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[feuille] if feuille in wb.sheetnames else wb.worksheets[0]
    lignes = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[4]:  # pas de n° dossier
            continue
        lignes.append({
            "numero_dossier": str(row[4]).strip(),
            "numero_client": str(row[2]).strip() if row[2] else "",
            "capital_rembourse": _f(row[6]),
            "interets_rembourses": _f(row[7]),
            "penalites_rembourses": _f(row[8]),
        })
    return lignes


def interets_par_hierarchie(remboursements, hierarchie_par_dossier):
    """Ventile les intérêts encaissés par agent, superviseur, agence.
    hierarchie_par_dossier : dict {numero_dossier: {'agent','superviseur','agence'}}
      → à construire depuis l'encours crédit du mois (fait_credit).
    Renvoie 3 dictionnaires + les non-rattachés (dossiers soldés/absents)."""
    par_agent = defaultdict(lambda: {"interets": 0.0, "capital": 0.0, "nb": 0})
    par_superviseur = defaultdict(lambda: {"interets": 0.0, "capital": 0.0, "nb": 0})
    par_agence = defaultdict(lambda: {"interets": 0.0, "capital": 0.0, "nb": 0})
    non_rattaches = {"interets": 0.0, "capital": 0.0, "nb": 0}

    for r in remboursements:
        h = hierarchie_par_dossier.get(r["numero_dossier"])
        it = r["interets_rembourses"]
        cap = r["capital_rembourse"]
        if h:
            for cle, dico in (("agent", par_agent), ("superviseur", par_superviseur),
                              ("agence", par_agence)):
                k = h.get(cle) or "(non renseigné)"
                dico[k]["interets"] += it
                dico[k]["capital"] += cap
                dico[k]["nb"] += 1
        else:
            non_rattaches["interets"] += it
            non_rattaches["capital"] += cap
            non_rattaches["nb"] += 1

    return {
        "par_agent": dict(par_agent),
        "par_superviseur": dict(par_superviseur),
        "par_agence": dict(par_agence),
        "non_rattaches": non_rattaches,
    }


def construire_hierarchie_depuis_credit(prets):
    """Depuis une liste de prêts (fait_credit du mois), construit
    {numero_dossier: {agent, superviseur, agence}}."""
    return {str(p.numero_dossier).strip(): {
        "agent": p.agent_credit, "superviseur": p.superviseur, "agence": p.agence
    } for p in prets}


if __name__ == "__main__":
    import sys
    rb = lire_remboursements(sys.argv[1])
    tot_int = sum(r["interets_rembourses"] for r in rb)
    print(f"Remboursements lus : {len(rb)} lignes | intérêts totaux : {tot_int:,.2f}")
    print("(La ventilation par agent nécessite l'encours crédit du même mois pour la jointure.)")
