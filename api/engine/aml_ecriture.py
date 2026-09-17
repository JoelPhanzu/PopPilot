"""Écriture du rapport AML/LBC-FT complet — remplit REPORTING LBC-FT depuis toutes les sources."""
from __future__ import annotations
import datetime as dt
import shutil
import os
import openpyxl
from engine.aml import operations_especes, transferts_grand_livre, portefeuille_client


def ecrire_aml(path_source_xls, sortie, periode_debut, periode_fin,
               path_inventaire=None, taux_cdf=None, db_path="socle/micropop.db",
               encours_credit_usd=None, nb_credits=None,
               credit_conso_cdf=None, nb_credits_conso=None, ligne_groupe=None):
    """Remplit le rapport LBC-FT.

    `ligne_groupe` : numero de la ligne « groupes » (statut juridique 4) dans la feuille
    « REPORTING LBC-FT ». A relever une fois sur le gabarit. Tant qu'il n'est pas fourni,
    les groupes sont comptes mais PAS ecrits, et le compte rendu le dit (`complet`).
    """
    # TAUX : lu dans param_taux_change à la date de fin de période (§42). AUCUN repli.
    #
    # Le repli précédent (`except: taux_cdf = 2263.57`) figeait le taux d'août 2026 :
    # tout rapport d'un autre mois sortait converti au mauvais taux, en silence, alors
    # que la doctrine est explicite — le taux est saisi, jamais figé. Si le taux du mois
    # n'est pas saisi, on refuse de produire le rapport et on dit quoi faire.
    if taux_cdf is None:
        from engine.etats_financiers import taux_change
        from socle.schema import get_session
        d = periode_fin.date() if hasattr(periode_fin, "date") else periode_fin
        s = get_session(db_path)
        try:
            taux_cdf = taux_change(s, d)      # lève un ValueError explicite si absent
        finally:
            s.close()

    # Dossier temporaire du système (« /tmp » n'existe pas sous Windows).
    import tempfile
    fd, tmp = tempfile.mkstemp(suffix=".xlsx", prefix="poppilot_aml_")
    os.close(fd)
    shutil.copy(path_source_xls, tmp)
    # La copie sert à la fois de source (brouillards, grand livre) et de classeur à
    # remplir : elle n'est supprimée qu'à la toute fin (elle contient des données clients).
    try:
        return _remplir_aml(tmp, sortie, periode_debut, periode_fin, path_inventaire,
                            taux_cdf, encours_credit_usd, nb_credits,
                            credit_conso_cdf, nb_credits_conso, ligne_groupe)
    finally:
        try:
            os.remove(tmp)
        except OSError:
            pass


def _remplir_aml(tmp, sortie, periode_debut, periode_fin, path_inventaire, taux_cdf,
                 encours_credit_usd, nb_credits, credit_conso_cdf, nb_credits_conso,
                 ligne_groupe=None):
    ops = operations_especes(tmp, taux_cdf)
    tr = transferts_grand_livre(tmp, taux_cdf)
    pf = portefeuille_client(path_inventaire, taux_cdf) if path_inventaire else None

    wb = openpyxl.load_workbook(tmp)
    ws = wb["REPORTING LBC-FT"]

    def sv(r, n, v):
        if n is not None: ws.cell(row=r, column=2, value=n)
        if v is not None: ws.cell(row=r, column=3, value=round(v, 2))

    # période
    ws.cell(row=5, column=2, value=periode_debut)
    ws.cell(row=6, column=2, value=periode_fin)

    # totaux espèces
    dep_nb, dep_cdf = ops["depot"]["total_nb"], ops["depot"]["total_cdf"]
    ret_nb, ret_cdf = ops["retrait"]["total_nb"], ops["retrait"]["total_cdf"]
    tot_esp_nb = dep_nb + ret_nb
    tot_esp_cdf = dep_cdf + ret_cdf
    tot_op_nb = tot_esp_nb + tr["nombre"]
    tot_op_cdf = tot_esp_cdf + tr["volume_cdf"]

    # 1. TAILLE
    sv(13, tot_op_nb, tot_op_cdf)          # Total Opérations = espèces + transferts
    sv(14, tot_esp_nb, tot_esp_cdf)        # Total Opérations en espèces = dépôts + retraits
    if encours_credit_usd is not None:
        sv(15, nb_credits, encours_credit_usd * taux_cdf)   # Total encours crédit (CDF)
    if credit_conso_cdf is not None:
        sv(16, nb_credits_conso, credit_conso_cdf)          # Crédits à la consommation
    sv(25, dep_nb, dep_cdf)                # Total Dépôts

    # 2. PORTEFEUILLE CLIENT
    #
    # LES GROUPES (statut juridique 4) SONT CALCULÉS MAIS N'ONT PAS ENCORE DE LIGNE.
    # `portefeuille_client` les compte bien (aout 2026 : 3 566 groupes, 732 M CDF) et le
    # dossier revendique « total clients = PP + PM + groupes ». Or seules les lignes PP
    # (35, 39) et PM (40) étaient écrites : les groupes étaient calculés puis JETÉS, et
    # le fichier remis à la Conformité SOUS-DÉCLARAIT le portefeuille, sans un mot.
    #
    # On ne devine pas la ligne : écrire des groupes dans la mauvaise case d'une
    # déclaration réglementaire est pire que l'omission. Tant que le numéro n'est pas
    # relevé sur le gabarit, on le passe par `ligne_groupe=` et, à défaut, le rapport
    # se DÉCLARE incomplet (clé `complet` / `non_ecrit` du compte rendu).
    non_ecrit = []
    if pf:
        sv(35, pf["pp_nombre"], pf["pp_solde"])   # Personnes physiques
        sv(39, pf["pp_nombre"], pf["pp_solde"])   # · Autres (= tout PP faute de sous-classif)
        sv(40, pf["pm_nombre"], pf["pm_solde"])   # Personnes morales
        if ligne_groupe:
            sv(ligne_groupe, pf["groupe_nombre"], pf["groupe_solde"])
        elif pf["groupe_nombre"]:
            non_ecrit.append(
                f"portefeuille GROUPE (statut 4) : {pf['groupe_nombre']} groupes, "
                f"{pf['groupe_solde']:,.2f} CDF — aucune ligne cible connue dans "
                f"« REPORTING LBC-FT ». Relever le numero de ligne sur le gabarit et le "
                f"passer en ligne_groupe= ; sans cela le portefeuille client declare est "
                f"incomplet.")

    # 3. Opérations en espèces par seuil
    sv(53, ops["depot"][">=10k"][0], ops["depot"][">=10k"][1])
    sv(54, ops["retrait"][">=10k"][0], ops["retrait"][">=10k"][1])
    sv(55, ops["depot"]["5k-10k"][0], ops["depot"]["5k-10k"][1])
    sv(56, ops["retrait"]["5k-10k"][0], ops["retrait"]["5k-10k"][1])

    # 7. LOCALISATION (provinces) — écrit à partir de L149
    # Même exigence qu'au portefeuille : une province calculée qui n'a pas de ligne
    # (« AUTRE », ou une agence absente de AGENCE_PROVINCE) ne doit pas disparaître en
    # silence. Elle n'est pas écrite — mais elle est DITE.
    if pf:
        prov_rows = {"HAUT-KATANGA": 149, "NORD-KIVU": 150, "KINSHASA": 151}
        for prov, v in pf["localisation"].items():
            r = prov_rows.get(prov)
            lib = str(ws.cell(row=r, column=1).value or "").upper() if r else ""
            # ne pas écraser un libellé existant : n'écrit que si la ligne correspond
            if r and (prov.split("-")[0] in lib or lib == ""):
                sv(r, v["nombre"], v["volume"])
            else:
                non_ecrit.append(
                    f"localisation « {prov} » : {v['nombre']} operations, "
                    f"{v['volume']:,.2f} CDF — aucune ligne cible dans la section 7.")
        if pf.get("agences_sans_province"):
            non_ecrit.append(
                "agences hors referentiel province (rangees en « AUTRE ») : "
                + ", ".join(pf["agences_sans_province"])
                + " — completer AGENCE_PROVINCE dans engine/aml.py.")

    wb.save(sortie)
    # `complet` répond avant l'envoi à la Conformité : tout ce qui est calculé a-t-il
    # été écrit ? Un rapport réglementaire incomplet doit se dénoncer lui-même.
    return {"sortie": sortie, "taux_cdf": taux_cdf,
            "total_operations": (tot_op_nb, tot_op_cdf),
            "total_especes": (tot_esp_nb, tot_esp_cdf),
            "total_depots": (dep_nb, dep_cdf),
            "transferts": tr, "portefeuille": pf,
            "depot_10k": ops["depot"][">=10k"], "depot_5k": ops["depot"]["5k-10k"],
            "retrait_5k": ops["retrait"]["5k-10k"],
            "complet": not non_ecrit, "non_ecrit": non_ecrit}
