"""
Tests : traitement SAGE (POST /sage/traiter), compte de résultat par agence (import),
primes hors « AC et SUP » (direction, support, recouvrement).

POURQUOI :
  - SAGE : un fichier équilibré mais converti au MAUVAIS taux s'importe sans erreur dans
    SAGE. On vérifie donc la conversion ligne à ligne au taux DU JOUR, le refus quand un
    jour n'a pas de taux, et le reformatage des comptes (exemples validés par le CDG).
  - Compte de résultat : il alimente les primes de direction. Un fichier incohérent
    (MICROPOP ≠ Σ agences) doit être refusé, un ré-import doit remplacer, pas doubler.
  - Primes : chiffres connus — JUNIOR 339,84 (fichier de primes), Victoire +38 316,10
    (compte de résultat de juillet), bornes du barème support.

Base SQLite de test uniquement : jamais Supabase (DATABASE_URL retirée après le chargement
de socle.schema, qui lit api/.env). Aucune donnée client : grand livre et tableau de
recouvrement fabriqués à la volée (noms d'agents anonymisés, montants du tableau CDG).

Lancer :  python tests/test_sage_primes.py
"""
import datetime as dt
import io
import os
import shutil
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import donnees_test as D   # noqa: E402

DB = "socle/test_sage_primes.db"
# Variantes du même fichier : NFC, NFD, et le nom abîmé par une décompression
# (« é » devenu « ├® », tel qu'il se trouve dans data_local).
NOMS_CR = ("COMPTE_RESULTAT_JUILLET_2026_isolé.xlsx",
           "COMPTE_RESULTAT_JUILLET_2026_isolé.xlsx",
           "COMPTE_RESULTAT_JUILLET_2026_isol├®.xlsx")
JUILLET = dt.date(2026, 7, 31)

_pret = False


def _preparer():
    global _pret
    import socle.schema as S
    os.environ.pop("DATABASE_URL", None)
    if _pret:
        return S
    S.fermer_moteurs()
    if os.path.exists(DB):
        os.remove(DB)
    S.init_db(DB)
    s = S.get_session(DB)
    for jour, taux in ((dt.date(2026, 9, 1), 2365.0), (dt.date(2026, 9, 2), 2370.0)):
        s.add(S.ParamTauxChange(date_effet=jour, devise_source="USD", devise_cible="CDF",
                                taux=taux))
    s.commit()
    s.close()
    import sage
    origine = S.get_session
    sage.get_session = lambda *a, **k: origine(DB)
    _pret = True
    return S


def _u(role):
    return {"login": role.lower(), "role": role,
            "agence": "AGENCE DE VICTOIRE" if role == "AGENCE" else None}


def _upload(contenu: bytes, nom: str):
    from fastapi import UploadFile
    return UploadFile(file=io.BytesIO(contenu), filename=nom)


def _xlsx(lignes) -> bytes:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    for l in lignes:
        ws.append(list(l))
    tampon = io.BytesIO()
    wb.save(tampon)
    return tampon.getvalue()


ENTETE_GL = ("Compte", "Libelle compte", "Sens", "Devise", "Montant Total",
             "Date Comptable", "Opération", "Utilisateur", "Guichet")
GL_OK = [
    ENTETE_GL,
    # exemple du README (validé CDG) : 5 114,47 USD × 2 365 = 12 095 721,55 CDF
    ("3.2.5.0.3", "Prets", "d", "USD", 5114.47, dt.datetime(2026, 9, 1), "op", "u", "g"),
    ("5.7.1.0.1", "Caisse USD", "c", "USD", 5114.47, dt.datetime(2026, 9, 1), "op", "u", "g"),
    # CDF : parité 1, suffixe 1 → 33114100 (exemple CDG)
    ("3.3.1.1.4", "Depots CDF", "c", "CDF", 100000, dt.datetime(2026, 9, 2), "op", "u", "g"),
    ("5.7.1.0.2", "Caisse CDF", "d", "CDF", 100000, dt.datetime(2026, 9, 2), "op", "u", "g"),
    # autre jour, AUTRE taux ; date en texte JJ/MM/AAAA comme certains exports CBS
    ("3.2.7.0.1.1", "Interets", "d", "USD", 10, "02/09/2026", "op", "u", "g"),
    ("5.7.1.0.1", "Caisse USD", "c", "USD", 10, "02/09/2026", "op", "u", "g"),
]


def _traiter(lignes, role="CDG"):
    import sage
    return sage.endpoint_traiter_sage(fichier=_upload(_xlsx(lignes), "GL_septembre.xlsx"),
                                      feuille=None, user=_u(role))


# ─────────────────────────────────────────────────────────────────────────────
def test_sage_conversion_taux_du_jour():
    _preparer()
    import openpyxl
    r = _traiter(GL_OK)
    assert r.headers["X-Sage-Statut"] == "OK", r.headers
    assert r.headers["X-Sage-Ecart"] == "0.00"
    ws = openpyxl.load_workbook(io.BytesIO(r.body)).active
    entete = [c.value for c in ws[1]]
    assert entete == ["Date Comptable", "N° Pièce", "Code journal", "CG", "Libelle compte",
                      "Devise", "Parité", "Montant devise", "Débit CDF", "Crédit CDF",
                      "N° Section", "Type_Ecriture"], entete
    l = {i: [c.value for c in ws[i]] for i in range(2, ws.max_row + 1)}
    assert l[2][3] == "32503000" and l[2][6] == 2365 and l[2][8] == 12095721.55, l[2]
    assert l[3][9] == 12095721.55 and l[3][8] is None
    assert l[4][3] == "33114100" and l[4][6] == 1 and l[4][9] == 100000, l[4]
    assert l[6][3] == "32701100" and l[6][6] == 2370 and l[6][8] == 23700, l[6]
    for ligne in l.values():
        assert ligne[1] is None and ligne[2] is None and ligne[10] is None   # vides
        assert ligne[11] == "G"


def test_sage_taux_manquant_refuse():
    _preparer()
    from fastapi import HTTPException
    lignes = GL_OK + [("3.2.5.0.3", "Prets", "d", "USD", 1, dt.datetime(2026, 9, 3),
                       "op", "u", "g")]
    try:
        _traiter(lignes)
        raise AssertionError("un jour sans taux a été converti")
    except HTTPException as e:
        assert e.status_code == 422 and "2026-09-03" in e.detail, e.detail


def test_sage_desequilibre_et_sens_inconnu_alertes():
    _preparer()
    r = _traiter(GL_OK + [("3.2.5.0.3", "Prets", "x", "USD", 1, dt.datetime(2026, 9, 1),
                           "op", "u", "g"),
                          ("3.2.5.0.3", "Prets", "d", "USD", 2, dt.datetime(2026, 9, 1),
                           "op", "u", "g")])
    assert r.headers["X-Sage-Statut"] == "ALERTE"
    assert r.headers["X-Sage-Ecart"] == f"{2 * 2365:.2f}"
    assert "1 ligne(s)" in r.headers["X-Sage-Alertes"]


def test_sage_roles_et_journal():
    S = _preparer()
    from fastapi import HTTPException
    for role in ("AUDIT", "AGENCE"):
        try:
            _traiter(GL_OK, role=role)
            raise AssertionError(f"{role} a lancé un traitement SAGE")
        except HTTPException as e:
            assert e.status_code == 403
    s = S.get_session(DB)
    statuts = [j.statut for j in s.query(S.JournalSageTraite).all()]
    s.close()
    assert "OK" in statuts and "ECHEC" in statuts and "ALERTE" in statuts, statuts


# ─────────────────────────────────────────────────────────────────────────────
def test_compte_resultat_agence_import_et_idempotence():
    S = _preparer()
    source = D.exiger_un_de(*NOMS_CR)
    from ingest.import_compte_resultat_agence import importer_compte_resultat_agence
    from engine.primes_categories import resultats_agences_depuis_base
    r1 = importer_compte_resultat_agence(source, JUILLET, db_path=DB)
    r2 = importer_compte_resultat_agence(source, JUILLET, db_path=DB)
    assert r2["remplacees"] == r1["lignes_importees"], (r1, r2)
    s = S.get_session(DB)
    n = s.query(S.CompteResultatAgence).count()
    s.close()
    assert n == r1["lignes_importees"], "le ré-import a dupliqué au lieu de remplacer"
    res = resultats_agences_depuis_base(JUILLET, db_path=DB)
    assert abs(res["AGENCE DE VICTOIRE"] - 38316.10) < 0.01, res
    assert abs(res["AGENCE OZONE"] - (-12498.93)) < 0.01
    assert abs(sum(res.values()) - 24798.80) < 0.01          # = colonne MICROPOP


def test_compte_resultat_incoherent_refuse():
    _preparer()
    import openpyxl
    source = D.exiger_un_de(*NOMS_CR)
    from ingest.import_compte_resultat_agence import importer_compte_resultat_agence
    dossier = tempfile.mkdtemp()
    try:
        copie = os.path.join(dossier, "cr_fausse.xlsx")
        shutil.copy(source, copie)
        wb = openpyxl.load_workbook(copie)
        ws = wb["Feuil2"]
        for row in ws.iter_rows():
            if row[0].value and "RESULTAT" in str(row[0].value).upper():
                row[7].value = 999999          # MICROPOP ≠ Σ agences
        wb.save(copie)
        try:
            importer_compte_resultat_agence(copie, dt.date(2026, 6, 30), db_path=DB)
            raise AssertionError("compte de résultat incohérent accepté")
        except ValueError as e:
            assert "incohérent" in str(e)
    finally:
        shutil.rmtree(dossier, ignore_errors=True)


def test_primes_direction_juillet():
    S = _preparer()
    D.exiger_un_de(*NOMS_CR)
    test_compte_resultat_agence_import_et_idempotence()
    import primes
    origine = S.get_session
    S.get_session = lambda *a, **k: origine(DB)
    try:
        r = primes.endpoint_primes_direction(arrete="2026-07-31", user=_u("CDG"))
    finally:
        S.get_session = origine
    ag = {l["agence"]: l for l in r["agences"]}
    assert ag["AGENCE DE VICTOIRE"]["prime_chef_agence"] == 383.16
    assert ag["AGENCE DE VICTOIRE"]["prime_adjoint"] == 191.58
    assert ag["AGENCE DE GOMBE"]["prime_chef_agence"] == 0          # perte
    dg = r["direction_generale"]["primes"]
    assert dg["Directeur Général"] == 247.99 and dg["Directeur Général Adjoint"] == 148.79
    assert dg["Directeur Administratif et Financier"] == 74.40
    assert dg["Responsable régional"] == 247.99
    from fastapi import HTTPException
    try:
        primes.endpoint_primes_direction(arrete="2026-07-31", user=_u("AGENCE"))
        raise AssertionError("une agence a lu les primes")
    except HTTPException as e:
        assert e.status_code == 403


# ─────────────────────────────────────────────────────────────────────────────
TABLEAU_RECOUVREMENT = [
    ("Équipe", "Agent", "Agence", "Montant 91-180", "Montant 181+", "Montant Radié"),
    ("Equipe A", "AGENT 01", "VICTOIRE", 0, 758, 176),
    ("Equipe A", "AGENT 02", "VICTOIRE", 2698, 0, 0),
    ("Equipe A", "AGENT 03", "LUBUMBASHI", 865.16, 9183, 1114),      # = JUNIOR : 339,84
    ("Equipe A", "AGENT 04", "OZONE", 0, 0, 240.01),
    ("Equipe A", "AGENT 05", "VICTOIRE", 0, 892.18, 0),
    ("Equipe A", "AGENT 06", "MASINA", 624.75, 177, 0),
    ("Equipe A", "AGENT 07", "MASINA", 0, 2397.65, 0),
    ("Equipe B", "AGENT 08", "OZONE", 0, 2816.03, 0),
    ("Equipe B", "AGENT 09", "GOMBE", 2208, 424, 0),
    ("Equipe B", "AGENT 10", "OZONE", 1813.44, 0, 0),
    ("Equipe B", "AGENT 11", "OZONE", 0, 0, 2048.34),
    ("TOTAL", None, None, "8 209,35", "16 647,86", "3 578,35"),      # texte, comme saisi
]


def test_primes_recouvrement():
    import primes
    r = primes.endpoint_primes_recouvrement(
        fichier=_upload(_xlsx(TABLEAU_RECOUVREMENT), "recouvrement.xlsx"), feuille=None,
        user=_u("CDG"))
    ag = {a["agent"]: a for a in r["agents"]}
    assert len(ag) == 11
    assert ag["AGENT 03"]["prime"] == 339.84, ag["AGENT 03"]
    assert ag["AGENT 01"]["prime"] == 31.54                   # 758×3 % + 176×5 %
    assert r["total_recouvre"] == {"91-180": 8209.35, "181-360": 16647.86, "radie": 3578.35}
    assert r["total_primes_agents"] == 760.45
    assert r["responsable"]["prime"] == 143.65                # 24,63 + 83,24 + 35,78
    assert r["alertes"] == []
    faux = TABLEAU_RECOUVREMENT[:-1] + [("TOTAL", None, None, 9000, 16647.86, 3578.35)]
    r = primes.endpoint_primes_recouvrement(
        fichier=_upload(_xlsx(faux), "recouvrement.xlsx"), feuille=None, user=_u("CDG"))
    assert r["alertes"] and "91-180" in r["alertes"][0]


def test_primes_support_bareme():
    from engine.primes_categories import primes_support
    base = {"taux_decaissement": 1.0, "epargne": 60, "encours": 100, "effectif": 4}
    cas = [(0.03, 30), (0.0301, 20), (0.05, 20), (0.0501, 10), (0.07, 10), (0.0701, 0)]
    for par30, prime_par in cas:
        l = primes_support([{**base, "agence": "A", "par30": par30}])["agences"][0]
        assert l["prime_par"] == prime_par, (par30, l)
        assert l["prime_unitaire"] == 5 + 10 + prime_par
        assert l["prime_totale_agence"] == (15 + prime_par) * 4
    l = primes_support([{**base, "agence": "A", "par30": 0.02, "taux_decaissement": 0.99,
                         "epargne": 59.9}])["agences"][0]
    assert (l["prime_decaissement"], l["prime_epargne"], l["prime_par"]) == (0, 0, 30)
    l = primes_support([{**base, "agence": "A", "par30": 0.02, "taux_decaissement": None,
                         "effectif": None}])["agences"][0]
    assert l["prime_totale_agence"] == 0 and not l["objectif_connu"]


# ─────────────────────────────────────────────────────────────────────────────
def _fichier_tmp(lignes, nom):
    chemin = os.path.join(tempfile.mkdtemp(), nom)
    with open(chemin, "wb") as f:
        f.write(_xlsx(lignes))
    return chemin


def test_import_taux_journaliers():
    S = _preparer()
    from ingest.import_taux import importer_taux
    f = _fichier_tmp([("Date", "Taux"), (dt.datetime(2026, 10, 1), 2300.5),
                      ("02/10/2026", "2 301,25"), (dt.date(2026, 10, 3), 2302)], "taux.xlsx")
    r = importer_taux(f, db_path=DB)
    assert (r["ajoutes"], r["inchanges"], r["du"], r["au"]) == (3, 0, "2026-10-01", "2026-10-03"), r
    assert importer_taux(f, db_path=DB)["inchanges"] == 3                  # rejouable
    s = S.get_session(DB)
    assert s.query(S.ParamTauxChange).filter_by(date_effet=dt.date(2026, 10, 2)).one().taux == 2301.25
    s.close()
    conflit = _fichier_tmp([("Date", "Taux"), ("01/10/2026", 9999)], "taux2.xlsx")
    try:
        importer_taux(conflit, db_path=DB)
        raise AssertionError("un taux existant a été écrasé sans accord")
    except ValueError as e:
        assert "01/10/2026" in str(e) and "Rien n'a été importé" in str(e)
    assert importer_taux(conflit, remplacer="oui", db_path=DB)["remplaces"] == 1
    for mauvais in ([("Date", "Taux"), ("01/13/2026", 2300)], [("Date", "Taux"), ("05/10/2026", 0)],
                    [("Jour", "Valeur"), ("05/10/2026", 2300)],
                    [("Date", "Taux"), ("06/10/2026", 2300), ("06/10/2026", 2310)]):
        try:
            importer_taux(_fichier_tmp(mauvais, "m.xlsx"), db_path=DB)
            raise AssertionError(f"fichier accepté : {mauvais}")
        except ValueError:
            pass


def test_journal_import_sans_date():
    """Le journal d'un import sans date ni exercice (taux) ne doit pas lever après coup."""
    S = _preparer()
    import import_cbs as I
    origine = S.get_session
    I.get_session = lambda *a, **k: origine(DB)
    try:
        I._tracer("taux_change", I.DOMAINES["taux_change"], {}, nom_fichier="taux.xlsx",
                  login="cdg", resultat={"acceptees": 3, "au": "2026-10-03"})
    finally:
        I.get_session = origine
    s = S.get_session(DB)
    l = s.query(S.ImportLog).filter_by(domaine="taux_change").one()
    s.close()
    assert l.date_arrete == dt.date(2026, 10, 3) and l.lignes_acceptees == 3


def test_primes_superviseurs_epargne():
    import primes
    lignes = [("Agence", "Cible", "Réalisation", "%"), ("VICTOIRE", 60500, 47162.76, 0.78),
              ("OZONE", 23500, 21212, 0.903), ("GOMBE", 5000, 6096.44, 1.219),
              ("MASINA", 48000, 33272, 0.693), ("LUBUMBASHI", 31000, 23238.44, 0.75),
              ("TOTAL", "168 000,00", "130 981,64", None)]
    _preparer()
    primes.BASE = DB
    r = primes.endpoint_primes_superviseurs_epargne(
        fichier=_upload(_xlsx(lignes), "epargne.xlsx"), mois="2026-05", feuille=None, user=_u("CDG"))
    ag = {a["agence"]: a for a in r["agences"]}
    assert round(ag["GOMBE"]["taux"] * 100, 1) == 121.9 and round(ag["VICTOIRE"]["taux"] * 100, 1) == 78.0
    # palier sur la réalisation TOTALE (aucune agence n'atteint 50 000) → 200 USD
    assert r["base_palier"] == "total" and r["date_arrete"] == "2026-05-31"
    assert r["total"]["realisation"] == 130981.64 and r["total"]["prime"] == 200
    assert r["alertes"] == []
    # conservé à son mois, relu sans fichier ; ré-import = remplacement, pas doublon
    r2 = primes.endpoint_primes_superviseurs_epargne(
        fichier=_upload(_xlsx(lignes), "epargne.xlsx"), mois="2026-05", feuille=None, user=_u("CDG"))
    assert r2["remplaces"] == 5
    g = primes.endpoint_collecte_epargne(mois="2026-05", user=_u("CDG"))
    assert g["total"]["realisation"] == 130981.64 and g["total"]["prime"] == 200
    assert len(g["agences"]) == 5


def test_compte_resultat_lecture_et_cloisonnement():
    S = _preparer()
    D.exiger_un_de(*NOMS_CR)
    test_compte_resultat_agence_import_et_idempotence()
    import compte_resultat as C
    origine = S.get_session
    C.get_session = lambda *a, **k: origine(DB)
    try:
        r = C.endpoint_compte_resultat(arrete="2026-07-31", user=_u("CDG"))
        v = C.endpoint_compte_resultat(arrete="2026-07-31", user=_u("AGENCE"))
        assert C.endpoint_arretes(user=_u("CDG"))["arretes"][0] == "2026-07-31"
    finally:
        C.get_session = origine
    assert len(r["agences"]) == 6
    p = {l["poste"]: l for l in r["postes"]}
    assert p["INTERETS SUR PRETS"]["nature"] == "produit"
    assert p["FRAIS BANCAIRES"]["nature"] == "charge"
    assert p["TOTAL PRODUITS"]["nature"] == "total_produits"
    res = next(l for l in r["postes"] if l["nature"] == "resultat")
    assert abs(res["montants"]["AGENCE DE VICTOIRE"] - 38316.10) < 0.01
    assert abs(res["total"] - 24798.80) < 0.01
    assert r["postes"][0]["poste"] == "INTERETS SUR PRETS"            # ordre du fichier
    assert v["agences"] == ["AGENCE DE VICTOIRE"]
    res_v = next(l for l in v["postes"] if l["nature"] == "resultat")
    assert set(res_v["montants"]) == {"AGENCE DE VICTOIRE"} and abs(res_v["total"] - 38316.10) < 0.01


if __name__ == "__main__":
    code = D.lancer("SAGE, compte de résultat agence, primes", [
        (test_sage_conversion_taux_du_jour, "SAGE : taux du jour, CG 8 chiffres, vides, G"),
        (test_sage_taux_manquant_refuse, "SAGE : jour sans taux -> 422 (jamais un taux voisin)"),
        (test_sage_desequilibre_et_sens_inconnu_alertes, "SAGE : déséquilibre / sens inconnu -> ALERTE"),
        (test_sage_roles_et_journal, "SAGE : AUDIT/AGENCE -> 403 ; journal_sage_traite"),
        (test_compte_resultat_agence_import_et_idempotence,
         "Compte de résultat juillet : Victoire 38 316,10 ; Σ = 24 798,80 ; ré-import"),
        (test_compte_resultat_incoherent_refuse, "Compte de résultat MICROPOP ≠ Σ agences -> refusé"),
        (test_primes_direction_juillet, "Primes direction juillet (Victoire 383,16 ; DG 247,99)"),
        (test_primes_recouvrement, "Primes recouvrement : 339,84 ; 760,45 ; resp. 143,65"),
        (test_primes_support_bareme, "Primes support : bornes du barème PAR, 5 $, 10 $"),
        (test_import_taux_journaliers, "Taux Date|Taux : ajout, rejouable, conflit refusé, remplacer=oui"),
        (test_journal_import_sans_date, "Journal d'un import sans date (taux) : pas d'erreur après coup"),
        (test_primes_superviseurs_epargne, "Superviseurs épargne : % recalculés, total 130 981,64"),
        (test_compte_resultat_lecture_et_cloisonnement,
         "Compte d'exploitation : ordre, natures, AGENCE limitée à sa colonne"),
    ])
    from socle.schema import fermer_moteurs
    fermer_moteurs()
    D.sortir(code)
