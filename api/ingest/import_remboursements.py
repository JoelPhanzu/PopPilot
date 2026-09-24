"""
Import des crédits remboursés (intérêts encaissés) → fait_remboursement_encaisse, avec la
hiérarchie (agent, superviseur, agence) RÉSOLUE à l'import par jointure sur le n° de dossier.

Sert à exposer la PRODUCTIVITÉ (intérêts, capital, pénalités encaissés par agent, superviseur,
agence). N'entre PAS dans le calcul des primes (décision CDG).

Lecture : colonnes du fichier CBS « Crédits remboursés » (feuille Worksheet) — Date de
remboursement | N° échéance | N° client | Noms | N° dossier | Montant déboursé | Capital |
Intérêts | Pénalités. Les montants sont contrôlés contre engine/interets_par_agent.lire_remboursements
(même total d'intérêts au centime, sinon refus) : une seule lecture fait foi.

RATTACHEMENT EN CASCADE (n° dossier) :
  1. encours du MOIS (fait_credit à l'arrêté) ;
  2. sinon encours de l'arrêté PRÉCÉDENT : un crédit soldé dans le mois a disparu de
     l'encours de fin de mois, mais ses intérêts ont bien été encaissés ce mois-ci ;
  3. sinon « non rattaché » (agent/superviseur/agence vides), compté et signalé.
Sur août 2026 : 1 211 remboursements non rattachés avec l'encours seul → 46 avec la cascade.
"""
from __future__ import annotations

import datetime as dt
import os

import openpyxl
from sqlalchemy import func, select

from socle.schema import FaitCredit, FaitRemboursementEncaisse, ImportLog, get_session, init_db
from engine.interets_par_agent import _f, lire_remboursements


def _date(v) -> dt.date | None:
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    try:
        return dt.datetime.strptime(str(v).strip(), "%d/%m/%Y").date()
    except (ValueError, TypeError):
        return None


def _hierarchie(s, date_arrete: dt.date) -> dict[str, tuple]:
    return {str(n).strip(): (a, x, g) for n, a, x, g in s.execute(
        select(FaitCredit.numero_dossier, FaitCredit.agent_credit, FaitCredit.superviseur,
               FaitCredit.agence).where(FaitCredit.date_arrete == date_arrete)).all()}


def importer_remboursements(path, date_arrete: dt.date, feuille: str | None = None,
                            db_path="socle/micropop.db") -> dict:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[feuille] if feuille and feuille in wb.sheetnames else wb.worksheets[0]
    lignes = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if len(row) < 9 or not row[4]:            # même règle que lire_remboursements
            continue
        lignes.append({
            "date_remboursement": _date(row[0]), "numero_echeance": str(row[1] or "").strip(),
            "numero_client": str(row[2]).strip() if row[2] else "",
            "numero_dossier": str(row[4]).strip(), "capital": _f(row[6]),
            "interets": _f(row[7]), "penalites": _f(row[8])})
    if not lignes:
        raise ValueError("Aucun remboursement lu : fichier « Crédits remboursés » attendu "
                         "(9 colonnes, n° de dossier en colonne E).")
    reference = lire_remboursements(path, feuille or "Worksheet")
    total, total_ref = sum(l["interets"] for l in lignes), sum(r["interets_rembourses"] for r in reference)
    if len(reference) != len(lignes) or abs(total - total_ref) > 0.005:
        raise ValueError(f"Lecture incohérente avec le moteur de référence ({len(lignes)} / "
                         f"{len(reference)} lignes, intérêts {total:,.2f} / {total_ref:,.2f}).")
    hors_mois = sum(1 for l in lignes if l["date_remboursement"] and
                    (l["date_remboursement"] > date_arrete or
                     l["date_remboursement"] < date_arrete.replace(day=1)))

    init_db(db_path)
    s = get_session(db_path)
    try:
        du_mois = _hierarchie(s, date_arrete)
        if not du_mois:
            raise ValueError(f"Aucun encours crédit à l'arrêté {date_arrete} : l'importer d'abord "
                             "(la hiérarchie agent/superviseur/agence vient de l'encours).")
        precedent = s.execute(select(func.max(FaitCredit.date_arrete))
                              .where(FaitCredit.date_arrete < date_arrete)).scalar()
        du_precedent = _hierarchie(s, precedent) if precedent else {}

        purges = s.query(FaitRemboursementEncaisse).filter(
            FaitRemboursementEncaisse.date_arrete == date_arrete).delete()
        sources = {"mois": 0, "precedent": 0, "non_rattache": 0}
        interets_non_rattaches = 0.0
        for l in lignes:
            h = du_mois.get(l["numero_dossier"])
            if h:
                sources["mois"] += 1
            else:
                h = du_precedent.get(l["numero_dossier"])
                if h:
                    sources["precedent"] += 1
                else:
                    sources["non_rattache"] += 1
                    interets_non_rattaches += l["interets"]
                    h = (None, None, None)
            s.add(FaitRemboursementEncaisse(
                date_arrete=date_arrete, date_remboursement=l["date_remboursement"],
                numero_dossier=l["numero_dossier"], numero_client=l["numero_client"],
                numero_echeance=l["numero_echeance"], capital_rembourse=l["capital"],
                interets_rembourses=l["interets"], penalites_rembourses=l["penalites"],
                agent_credit=h[0], superviseur=h[1], agence=h[2]))
        resultat = {"lignes_importees": len(lignes), "remplacees": purges,
                    "interets_total": round(total, 2),
                    "rattaches_encours_du_mois": sources["mois"],
                    "rattaches_encours_precedent": sources["precedent"],
                    "encours_precedent": precedent.isoformat() if precedent else None,
                    "non_rattaches": sources["non_rattache"],
                    "interets_non_rattaches": round(interets_non_rattaches, 2),
                    "dates_hors_du_mois": hors_mois}
        s.add(ImportLog(domaine="remboursements", fichier=os.path.basename(str(path)),
                        date_snapshot=dt.date.today(), date_arrete=date_arrete,
                        lignes_acceptees=len(lignes), lignes_rejetees=0,
                        horodatage=dt.datetime.now(), message=str(resultat)[:500]))
        s.commit()
    finally:
        s.close()
    return resultat
