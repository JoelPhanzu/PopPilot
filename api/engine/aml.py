"""Moteur AML / LBC-FT (Rapport 5) — version complète. CLAUDE.md §52-58."""
from __future__ import annotations
import openpyxl

AGENCE_PROVINCE = {
    "AGENCE DE VICTOIRE": "KINSHASA", "AGENCE OZONE": "KINSHASA",
    "AGENCE DE MASINA": "KINSHASA", "AGENCE DE GOMBE": "KINSHASA",
    "AGENCE DE GOMA": "NORD-KIVU", "AGENCE DE LUBUMBASHI": "HAUT-KATANGA",
}

def _f(v):
    if v in (None, ""): return 0.0
    try: return float(str(v).replace(",", ".").replace("\xa0", "").replace(" ", ""))
    except (ValueError, TypeError): return 0.0

def _est_depot(l):
    l = (l or "").lower()
    return ("dépôt" in l or "d‚p" in l or "depot" in l) or ("pargne à la carte" in l)

def _est_retrait(l):
    return "retrait" in (l or "").lower()

def _seuil(usd):
    if usd >= 10000: return ">=10k"
    if usd >= 5000: return "5k-10k"
    return None

def operations_especes(path_xlsx, taux_cdf):
    wb = openpyxl.load_workbook(path_xlsx, read_only=True, data_only=True)
    res = {"depot": {">=10k": [0, 0.0], "5k-10k": [0, 0.0], "total_nb": 0, "total_cdf": 0.0},
           "retrait": {">=10k": [0, 0.0], "5k-10k": [0, 0.0], "total_nb": 0, "total_cdf": 0.0}}
    for feuille, devise in [("Brouillard de caisse USD", "USD"), ("Brouillard de caisse CDF", "CDF")]:
        if feuille not in wb.sheetnames: continue
        ws = wb[feuille]
        for row in ws.iter_rows(min_row=4, values_only=True):
            lib = str(row[4] or ""); deb = _f(row[7]); cred = _f(row[8])
            if _est_depot(lib) and deb > 0:
                usd = deb if devise == "USD" else deb / taux_cdf
                cdf = deb * taux_cdf if devise == "USD" else deb
                res["depot"]["total_nb"] += 1; res["depot"]["total_cdf"] += cdf
                s = _seuil(usd)
                if s: res["depot"][s][0] += 1; res["depot"][s][1] += usd
            elif _est_retrait(lib) and cred > 0:
                usd = cred if devise == "USD" else cred / taux_cdf
                cdf = cred * taux_cdf if devise == "USD" else cred
                res["retrait"]["total_nb"] += 1; res["retrait"]["total_cdf"] += cdf
                s = _seuil(usd)
                if s: res["retrait"][s][0] += 1; res["retrait"][s][1] += usd
    return res

def transferts_grand_livre(path_xlsx, taux_cdf=1.0):
    """Transferts (GL 330/331/332, libellé transfert). Montant = débit OU crédit (col 5/6).
    RÈGLE (CDG) : le grand livre est TOUJOURS en USD → conversion systématique en CDF au taux."""
    wb = openpyxl.load_workbook(path_xlsx, read_only=True, data_only=True)
    if "Grand Livre" not in wb.sheetnames:
        return {"nombre": 0, "volume_cdf": 0.0, "volume_usd": 0.0}
    ws = wb["Grand Livre"]
    nb = 0; vol_usd = 0.0
    for row in ws.iter_rows(min_row=2, values_only=True):
        compte = str(row[0] or "")
        lib = str(row[3] or "").lower()
        if compte[:3] in ("330", "331", "332") and "transf" in lib:
            # montant = débit (col 5) si rempli, sinon crédit (col 6). GL = USD (toujours).
            montant = _f(row[5]) if row[5] not in (None, "") else _f(row[6])
            nb += 1
            vol_usd += abs(montant)
    return {"nombre": nb, "volume_usd": vol_usd, "volume_cdf": vol_usd * taux_cdf}

def portefeuille_client(path_inventaire, taux_cdf=2263.57):
    """Portefeuille client depuis l'inventaire dépôt. Soldes en CDF (USD converti au taux).
    Statut juridique : 1 = personne physique, 2 = personne morale, 4 = groupe (INCLUS).
    Solde = colonne solde_fin (index 23). Devise par compte (col 4) : USD → CDF."""
    wb = openpyxl.load_workbook(path_inventaire, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    pp_clients, pm_clients, grp_clients = set(), set(), set()
    pp_solde = pm_solde = grp_solde = 0.0
    loc = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        statut = str(row[17] or "").strip()
        id_client = row[5]
        devise = str(row[4] or "").strip().upper()
        solde_fin = _f(row[23])                          # colonne solde_fin
        # conversion en CDF selon la devise du compte
        solde_cdf = solde_fin * taux_cdf if devise == "USD" else solde_fin
        agence = str(row[24] or "").strip().upper()
        prov = AGENCE_PROVINCE.get(agence, "AUTRE")
        if statut == "1":
            pp_clients.add(id_client); pp_solde += solde_cdf
        elif statut == "2":
            pm_clients.add(id_client); pm_solde += solde_cdf
        elif statut == "4":
            grp_clients.add(id_client); grp_solde += solde_cdf
        # Localisation = OPÉRATIONS (dépôts + retraits) par province, en CDF.
        # Chaque mouvement > 0 = 1 opération (dépôt ET retrait sur la ligne = 2 opérations).
        depot = _f(row[21])      # montant_depot
        retrait = _f(row[22])    # montant_retrait
        if devise == "USD":
            depot *= taux_cdf; retrait *= taux_cdf
        l = loc.setdefault(prov, [0, 0.0])
        if depot > 0:
            l[0] += 1; l[1] += depot
        if retrait > 0:
            l[0] += 1; l[1] += retrait
    return {
        "pp_nombre": len(pp_clients), "pp_solde": pp_solde,
        "pm_nombre": len(pm_clients), "pm_solde": pm_solde,
        "groupe_nombre": len(grp_clients), "groupe_solde": grp_solde,
        "clients_total": len(pp_clients) + len(pm_clients) + len(grp_clients),
        "solde_total": pp_solde + pm_solde + grp_solde,
        "localisation": {k: {"nombre": v[0], "volume": v[1]} for k, v in loc.items()},
    }
