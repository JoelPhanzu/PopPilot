"""
Import du roster mensuel + objectifs (Lot 1.6 amont) — CLAUDE.md §5, §15.3, §18.2.

Le fichier OBJECTIF est DOUBLE source :
  - dim_employe : la liste officielle des agents/superviseurs actifs (→ détection orphelins).
  - param_objectif : leurs objectifs (décaissement nombre, volume, portefeuille, PAR).
Versionné à date d'effet : le roster de mai ne s'applique qu'aux calculs de mai (§18.4).
"""
from __future__ import annotations

import datetime as dt

import openpyxl

from socle.schema import DimEmploye, ParamObjectif, get_session, init_db
from socle import historisation as _noop  # noqa


def _f(v):
    if v in (None, ""):
        return 0.0
    try:
        return float(str(v).replace(",", ".").replace("\xa0", "").replace(" ", ""))
    except ValueError:
        return 0.0


def _par(v):
    """PAR objectif en fraction : « 5% » → 0,05 ; 5 → 0,05 ; 0,05 → 0,05."""
    if v in (None, ""):
        return 0.0
    texte = str(v).strip()
    x = _f(texte.rstrip("%").strip())
    return x / 100 if texte.endswith("%") or x > 1 else x


def _sans_accent(t) -> str:
    import unicodedata
    t = unicodedata.normalize("NFKD", str(t or ""))
    return "".join(c for c in t if not unicodedata.combining(c)).upper()


def colonnes(entete) -> dict | None:
    """Repère les colonnes par leur EN-TÊTE, quel que soit le format :
    7 colonnes (AGENCE | SUPERVISEUR | AGENT DE CREDIT | #NOMBRE A DECAISSE | VOLUME |
    PORTEFEUILLE | PAR) ou 8 (PORTEFEUILLE volume + PORTEFEUILLE nombre de clients).
    None si la ligne n'est pas un en-tête."""
    e = [_sans_accent(c) for c in entete]
    trouve = lambda pred: next((i for i, c in enumerate(e) if pred(c)), None)  # noqa: E731
    agence = trouve(lambda c: c.strip() == "AGENCE" or c.startswith("AGENCE"))
    agent = trouve(lambda c: "AGENT" in c)
    if agence is None or agent is None:
        return None
    portef = [i for i, c in enumerate(e) if "PORTEFEUILLE" in c]
    vol_portef = next((i for i in portef if "CLIENT" not in e[i] and "NOMBRE" not in e[i]), None)
    return {
        "agence": agence, "agent": agent,
        "superviseur": trouve(lambda c: "SUPERVISEUR" in c),
        "nombre": trouve(lambda c: "NOMBRE" in c and "PORTEFEUILLE" not in c and "CLIENT" not in c),
        "volume": trouve(lambda c: c.strip().startswith("VOLUME")),
        "portefeuille": vol_portef,
        "par": trouve(lambda c: c.strip() == "PAR" or c.strip().startswith("PAR ")),
    }


# Format historique sans en-tête reconnaissable (positions A→H).
_POSITIONS = {"agence": 0, "superviseur": 1, "agent": 2, "nombre": 3, "volume": 4,
              "portefeuille": 5, "par": 7}


def importer_objectifs(path, date_effet: dt.date, db_path="socle/micropop.db"):
    init_db(db_path)
    s = get_session(db_path)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["OBJECTIF"] if "OBJECTIF" in wb.sheetnames else wb.worksheets[0]
    lignes = list(ws.iter_rows(values_only=True))
    wb.close()

    # En-tête : première ligne (parmi les 10 premières) où l'on reconnaît AGENCE et AGENT.
    cols, debut = None, 1
    for i, row in enumerate(lignes[:10]):
        cols = colonnes(row)
        if cols:
            debut = i + 1
            break
    cols = cols or _POSITIONS
    val = lambda row, cle: (row[cols[cle]] if cols.get(cle) is not None  # noqa: E731
                            and cols[cle] < len(row) else None)

    # purge du roster + objectifs de cette date d'effet (idempotence)
    s.query(DimEmploye).filter(DimEmploye.date_debut == date_effet).delete()
    s.query(ParamObjectif).filter(ParamObjectif.date_effet == date_effet).delete()

    agents = superviseurs = 0
    superviseurs_vus = set()
    for row in lignes[debut:]:
        agence, superviseur, agent = val(row, "agence"), val(row, "superviseur"), val(row, "agent")
        if not agent or not str(agent).strip():
            continue
        agence = " ".join(str(agence or "").split())
        agent = " ".join(str(agent).split())
        superviseur = " ".join(str(superviseur).split()) if superviseur else ""

        # agent de crédit
        s.add(DimEmploye(nom=agent, fonction="agent_credit", agence=agence,
                         date_debut=date_effet, date_fin=None))
        agents += 1
        # superviseur (une fois par nom+agence)
        if superviseur and (superviseur, agence) not in superviseurs_vus:
            s.add(DimEmploye(nom=superviseur, fonction="superviseur", agence=agence,
                             date_debut=date_effet, date_fin=None))
            superviseurs_vus.add((superviseur, agence))
            superviseurs += 1
        # objectifs de l'agent (cellule vide = 0 : agent au roster sans objectif)
        s.add(ParamObjectif(
            date_effet=date_effet, agence=agence, agent=agent,
            objectif_decaissement_nombre=_f(val(row, "nombre")),   # #NOMBRE A DECAISSE
            objectif_volume=_f(val(row, "volume")),                # VOLUME
            objectif_portefeuille=_f(val(row, "portefeuille")),    # PORTEFEUILLE (volume)
            objectif_par=_par(val(row, "par")),                    # PAR
        ))
    s.commit()
    n_obj = s.query(ParamObjectif).filter(ParamObjectif.date_effet == date_effet).count()
    s.close()
    return {"agents": agents, "superviseurs": superviseurs, "objectifs": n_obj}


if __name__ == "__main__":
    import sys
    path = sys.argv[1]
    eff = dt.date.fromisoformat(sys.argv[2]) if len(sys.argv) > 2 else dt.date(2026, 5, 1)
    print(importer_objectifs(path, eff))
