"""
Tests du module Archives : bibliothèque versionnée, édition en ligne tracée, séries.

POURQUOI :
  - Remplacer un rapport ne doit JAMAIS perdre l'ancien (historique = preuve).
  - L'édition en ligne ne doit jamais toucher le fichier déposé, et chaque modification
    doit rester tracée (qui, quand) même quand une cellule est modifiée deux fois.
  - Les séries « calcul_poppilot » doivent redonner les chiffres des moteurs au centime
    (mai : encours 10 814 330,66 ; PAR30 1 052 118,05 ; provisions 938 244,42), et un
    historique importé ne doit pas écraser un point recalculé.

Base SQLite + dossier d'archives temporaires : jamais Supabase, jamais le vrai dépôt.

Lancer :  python tests/test_archives.py
"""
import datetime as dt
import io
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import donnees_test as D   # noqa: E402

DB = "socle/test_archives.db"
MAI = dt.date(2026, 5, 30)
_etat = {}


def _preparer(avec_credit=False):
    import socle.schema as S
    import archives as A
    if not _etat:
        os.environ.pop("DATABASE_URL", None)
        os.environ["POPPILOT_ARCHIVES_DIR"] = tempfile.mkdtemp(prefix="archives_test_")
        S.fermer_moteurs()
        if os.path.exists(DB):
            os.remove(DB)
        from socle.seed_parametres import seed
        seed(DB)
        origine = S.get_session
        A.get_session = lambda *a, **k: origine(DB)
        _etat["ok"] = True
    if avec_credit and "credit" not in _etat:
        from ingest.import_credit import importer_credit
        importer_credit(D.exiger("Enours_MAI_2026_.xls"), MAI, date_snapshot=dt.date(2026, 6, 1),
                        db_path=DB)
        _etat["credit"] = True
    return S, A


def _u(role):
    return {"login": role.lower(), "role": role,
            "agence": "AGENCE DE VICTOIRE" if role == "AGENCE" else None}


def _xlsx(lignes) -> bytes:
    import openpyxl
    wb = openpyxl.Workbook()
    for l in lignes:
        wb.active.append(list(l))
    t = io.BytesIO()
    wb.save(t)
    return t.getvalue()


def _upload(contenu, nom):
    from fastapi import UploadFile
    return UploadFile(file=io.BytesIO(contenu), filename=nom)


ORIGINAL = _xlsx([("Agence", "Encours"), ("VICTOIRE", 100), ("OZONE", 200)])


def _deposer(A, role="CDG", remplace_id=None, contenu=ORIGINAL, nom="portee_2025.xlsx"):
    return A.deposer(fichier=_upload(contenu, nom), titre="Rapport de portée 2025",
                     type_rapport="portee", periode="2025", remplace_id=remplace_id, user=_u(role))


# ─────────────────────────────────────────────────────────────────────────────
def test_depot_versions_et_roles():
    S, A = _preparer()
    from fastapi import HTTPException
    v1 = _deposer(A)
    assert v1["version"] == 1 and v1["editable"]
    v2 = _deposer(A, remplace_id=v1["id"], contenu=_xlsx([("Agence", "Encours"), ("VICTOIRE", 150)]))
    assert v2["version"] == 2 and v2["remplace_id"] == v1["id"]
    ids = [a["id"] for a in A.lister(user=_u("AUDIT"))["archives"]]
    assert v2["id"] in ids and v1["id"] not in ids                 # seule la version courante
    chaine = [v["version"] for v in A.versions(v2["id"], user=_u("CDG"))["versions"]]
    assert chaine == [2, 1]
    fichier = A.telecharger(v1["id"], user=_u("CDG"))              # v1 toujours téléchargeable
    assert open(fichier.path, "rb").read() == ORIGINAL
    for appel, code in ((lambda: _deposer(A, remplace_id=v1["id"]), 409),
                        (lambda: _deposer(A, role="AUDIT"), 403),
                        (lambda: A.lister(user=_u("AGENCE")), 403)):
        try:
            appel()
            raise AssertionError(f"attendu {code}")
        except HTTPException as e:
            assert e.status_code == code, (e.status_code, code)


def test_edition_en_ligne_tracee():
    S, A = _preparer()
    a = _deposer(A)
    d = A.lire_donnees(a["id"], user=_u("CDG"))
    assert d["grille"] == [["Agence", "Encours"], ["VICTOIRE", "100"], ["OZONE", "200"]]
    M = A.Modifications
    A.modifier_donnees(a["id"], M(modifications=[
        {"ligne": 1, "colonne": 1, "valeur": "110"},
        {"ligne": 3, "colonne": 0, "valeur": "GOMBE"},             # nouvelle ligne
        {"ligne": 0, "colonne": 2, "valeur": "PAR30"}]),            # nouvelle colonne
        user=_u("CDG"))
    A.modifier_donnees(a["id"], M(modifications=[{"ligne": 1, "colonne": 1, "valeur": "120"}]),
                       user=_u("DIRECTION"))
    d = A.lire_donnees(a["id"], user=_u("AUDIT"))
    assert d["grille"] == [["Agence", "Encours", "PAR30"], ["VICTOIRE", "120", ""],
                           ["OZONE", "200", ""], ["GOMBE", "", ""]], d["grille"]
    assert d["nb_modifications"] == 4 and d["derniere_modification"]["par"] == "direction"
    assert open(A.telecharger(a["id"], user=_u("CDG")).path, "rb").read() == ORIGINAL   # intact
    import openpyxl
    ws = openpyxl.load_workbook(io.BytesIO(A.exporter(a["id"], user=_u("CDG")).body)).active
    assert ws["B2"].value == 120 and ws["A4"].value == "GOMBE"      # nombre redevenu nombre
    from fastapi import HTTPException
    try:
        A.modifier_donnees(a["id"], M(modifications=[{"ligne": 1, "colonne": 1, "valeur": "x"}]),
                           user=_u("AUDIT"))
        raise AssertionError("l'AUDIT a modifié une archive")
    except HTTPException as e:
        assert e.status_code == 403
    pdf = A.deposer(fichier=_upload(b"%PDF-1.4 test", "rapport.pdf"), titre="PDF", type_rapport="autre",
                    periode="", remplace_id=None, user=_u("CDG"))
    try:
        A.lire_donnees(pdf["id"], user=_u("CDG"))
        raise AssertionError("un PDF s'est ouvert en édition")
    except HTTPException as e:
        assert e.status_code == 422


def test_series_moteurs_et_historique():
    _preparer(avec_credit=True)
    from engine.series import alimenter_depuis_moteurs, arretes_non_alimentes, importer_series
    assert "2026-05-30" in arretes_non_alimentes(db_path=DB)
    r1 = alimenter_depuis_moteurs(MAI, db_path=DB)
    r2 = alimenter_depuis_moteurs(MAI, db_path=DB)
    assert r1["ajout"] == r1["points"] and r2["ajout"] == 0 and r2["maj"] == r1["points"]
    assert "2026-05-30" not in arretes_non_alimentes(db_path=DB)
    from socle.schema import get_session, SerieIndicateur as SI
    s = get_session(DB)

    def v(ind, agence=None):
        q = s.query(SI).filter(SI.indicateur == ind, SI.date_arrete == MAI)
        q = q.filter(SI.agence.is_(None)) if agence is None else q.filter(SI.agence == agence)
        return q.one().valeur
    try:
        assert abs(v("encours_credit") - 10814330.66) < 0.01
        assert abs(v("par30") - 1052118.05) < 0.01
        assert abs(v("provisions") - 938244.42) < 0.01
        assert abs(v("encours_credit", "AGENCE DE VICTOIRE") - 2531178.47) < 0.01
    finally:
        s.close()
    f = os.path.join(tempfile.mkdtemp(), "hist.xlsx")
    with open(f, "wb") as h:
        h.write(_xlsx([("Indicateur", "Date", "Agence", "Valeur", "Unité"),
                       ("encours_credit", "31/12/2023", None, "8 500 000,50", "USD"),
                       ("encours_credit", dt.datetime(2026, 5, 30), None, 1, "USD")]))
    r = importer_series(f, db_path=DB)
    assert (r["ajout"], r["garde"]) == (1, 1), r                    # le calcul prime sur l'import
    s = get_session(DB)
    try:
        assert abs(v("encours_credit") - 10814330.66) < 0.01
        assert s.query(SI).filter(SI.date_arrete == dt.date(2023, 12, 31)).one().valeur == 8500000.5
    finally:
        s.close()


if __name__ == "__main__":
    code = D.lancer("Archives (bibliothèque, édition, séries)", [
        (test_depot_versions_et_roles, "Dépôt, v2 garde v1, chaîne des versions, rôles"),
        (test_edition_en_ligne_tracee, "Édition : lignes/colonnes ajoutées, trace, fichier intact"),
        (test_series_moteurs_et_historique, "Séries mai = moteurs au centime ; import ne l'écrase pas"),
    ])
    from socle.schema import fermer_moteurs
    fermer_moteurs()
    D.sortir(code)
