"""
PopPilot API — module ARCHIVES (étape 6) : bibliothèque, édition en ligne, séries.

BIBLIOTHÈQUE (archive_rapport)
  GET  /archives                      dernière version de chaque rapport (filtres type/période)
  POST /archives        (multipart)   dépôt ; avec remplace_id → NOUVELLE VERSION (l'ancienne reste)
  GET  /archives/{id}/versions        chaîne des versions, de la plus récente à la plus ancienne
  GET  /archives/{id}/fichier         téléchargement du fichier déposé (jamais modifié)

ÉDITION EN LIGNE (archive_donnees) — tableurs .xlsx / .csv
  GET  /archives/{id}/donnees         grille = fichier d'origine + modifications
  POST /archives/{id}/donnees         modifications {ligne, colonne, valeur} (ajout de lignes /
                                      colonnes = cellules au-delà de la grille)
  GET  /archives/{id}/export          grille modifiée en .xlsx
  Chaque modification est AJOUTÉE (modifie_par, modifie_le) : l'historique n'est jamais écrasé,
  la valeur affichée est la plus récente. Le fichier déposé reste intact.

SÉRIES (serie_indicateur)
  GET  /series/indicateurs            indicateurs disponibles (+ arrêtés crédit non alimentés)
  GET  /series?indicateur=&agence=    points de la série
  POST /series/import   (multipart)   historiques Excel (Indicateur | Date | Agence | Valeur | Unité)
  POST /series/alimenter?arrete=      calcul par les moteurs validés pour un arrêté chargé

STOCKAGE : les fichiers vont dans POPPILOT_ARCHIVES_DIR (défaut api/archives_depot/, hors
dépôt git) — sur le serveur qui héberge l'API. fichier_url = chemin RELATIF à ce dossier.

RÔLES : lecture = accès total (DIRECTION, CDG, AUDIT) ; dépôt, édition, import, calcul =
DIRECTION / CDG. Les archives sont institutionnelles : un rôle AGENCE n'y a pas accès.
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import os
import uuid

import openpyxl
from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse, Response
from pydantic import BaseModel
from sqlalchemy import func, select

from socle.schema import ArchiveDonnees, ArchiveRapport, SerieIndicateur, get_session
from engine.series import alimenter_depuis_moteurs, arretes_non_alimentes, importer_series

from auth_supabase import ROLES_ACCES_TOTAL, ROLES_ECRITURE, exiger_role, utilisateur_courant
from import_cbs import _ecrire_sur_disque, _nom_sain

routeur = APIRouter(tags=["archives"])

EXTENSIONS = (".xlsx", ".xlsm", ".xls", ".csv", ".pdf", ".docx", ".pptx")
EDITABLES = (".xlsx", ".xlsm", ".csv")
MAX_LIGNES, MAX_COLONNES = 5000, 100
TYPE_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def dossier_archives() -> str:
    d = os.environ.get("POPPILOT_ARCHIVES_DIR") or os.path.join(
        os.path.dirname(os.path.abspath(__file__)), "archives_depot")
    os.makedirs(d, exist_ok=True)
    return d


def _archive(s, archive_id: int) -> ArchiveRapport:
    a = s.get(ArchiveRapport, archive_id)
    if a is None:
        raise HTTPException(404, f"Archive {archive_id} introuvable.")
    return a


def _vue(a: ArchiveRapport) -> dict:
    return {"id": a.id, "titre": a.titre, "type_rapport": a.type_rapport, "periode": a.periode,
            "format": a.format, "version": a.version, "depose_par": a.depose_par,
            "date_depot": a.date_depot.isoformat() if a.date_depot else None,
            "remplace_id": a.remplace_id, "editable": f".{a.format}" in EDITABLES}


# ─── Bibliothèque ────────────────────────────────────────────────────────────
@routeur.get("/archives")
def lister(type_rapport: str | None = None, periode: str | None = None,
           user: dict = Depends(utilisateur_courant)):
    exiger_role(user, ROLES_ACCES_TOTAL)
    s = get_session()
    try:
        remplaces = set(s.execute(select(ArchiveRapport.remplace_id)
                                  .where(ArchiveRapport.remplace_id.is_not(None))).scalars())
        q = select(ArchiveRapport).order_by(ArchiveRapport.date_depot.desc())
        if type_rapport:
            q = q.where(ArchiveRapport.type_rapport == type_rapport)
        if periode:
            q = q.where(ArchiveRapport.periode == periode)
        courantes = [a for a in s.execute(q).scalars() if a.id not in remplaces]
        types = sorted(t for t in s.execute(select(ArchiveRapport.type_rapport).distinct()).scalars() if t)
    finally:
        s.close()
    return {"archives": [_vue(a) for a in courantes], "types": types}


@routeur.post("/archives")
def deposer(fichier: UploadFile = File(...), titre: str = Form(...),
            type_rapport: str = Form("autre"), periode: str = Form(""),
            remplace_id: int | None = Form(None), user: dict = Depends(utilisateur_courant)):
    exiger_role(user, ROLES_ECRITURE)
    nom = _nom_sain(fichier.filename, EXTENSIONS)
    extension = os.path.splitext(nom)[1].lower()
    relatif = f"{dt.date.today():%Y/%m}/{uuid.uuid4().hex}{extension}"
    chemin = os.path.join(dossier_archives(), *relatif.split("/"))
    os.makedirs(os.path.dirname(chemin), exist_ok=True)
    try:
        _ecrire_sur_disque(fichier, chemin)
    except HTTPException:
        if os.path.exists(chemin):
            os.remove(chemin)
        raise
    s = get_session()
    try:
        version = 1
        if remplace_id is not None:
            ancienne = _archive(s, remplace_id)
            deja = s.execute(select(ArchiveRapport.id).where(
                ArchiveRapport.remplace_id == remplace_id)).first()
            if deja:
                raise HTTPException(409, f"L'archive {remplace_id} a déjà été remplacée "
                                         f"(par {deja[0]}) : remplacer la version la plus récente.")
            version = (ancienne.version or 1) + 1
        a = ArchiveRapport(titre=titre.strip()[:200], type_rapport=type_rapport.strip() or "autre",
                           periode=periode.strip() or None, fichier_url=relatif,
                           format=extension.lstrip("."), version=version,
                           depose_par=user["login"], date_depot=dt.datetime.now(),
                           remplace_id=remplace_id)
        s.add(a)
        s.commit()
        return _vue(a)
    except Exception:
        s.rollback()
        if os.path.exists(chemin):
            os.remove(chemin)
        raise
    finally:
        s.close()


@routeur.get("/archives/{archive_id}/versions")
def versions(archive_id: int, user: dict = Depends(utilisateur_courant)):
    exiger_role(user, ROLES_ACCES_TOTAL)
    s = get_session()
    try:
        chaine, a = [], _archive(s, archive_id)
        while a is not None:
            chaine.append(_vue(a))
            a = s.get(ArchiveRapport, a.remplace_id) if a.remplace_id else None
    finally:
        s.close()
    return {"versions": chaine}


@routeur.get("/archives/{archive_id}/fichier")
def telecharger(archive_id: int, user: dict = Depends(utilisateur_courant)):
    exiger_role(user, ROLES_ACCES_TOTAL)
    s = get_session()
    try:
        a = _archive(s, archive_id)
        relatif, titre, fmt, version = a.fichier_url, a.titre, a.format, a.version
    finally:
        s.close()
    chemin = os.path.realpath(os.path.join(dossier_archives(), *relatif.split("/")))
    if not chemin.startswith(os.path.realpath(dossier_archives())) or not os.path.exists(chemin):
        raise HTTPException(410, "Fichier absent du dépôt d'archives (déplacé ou supprimé du serveur).")
    nom = "".join(c if c.isalnum() or c in "-_ " else "_" for c in titre)[:80]
    return FileResponse(chemin, filename=f"{nom}_v{version}.{fmt}")


# ─── Édition en ligne ────────────────────────────────────────────────────────
def _grille_origine(a: ArchiveRapport) -> list[list[str]]:
    chemin = os.path.join(dossier_archives(), *a.fichier_url.split("/"))
    if a.format == "csv":
        with open(chemin, encoding="utf-8-sig", newline="") as f:
            echantillon = f.read(4096)
            f.seek(0)
            sep = ";" if echantillon.count(";") > echantillon.count(",") else ","
            return [[c for c in l[:MAX_COLONNES]] for _, l in zip(range(MAX_LIGNES), csv.reader(f, delimiter=sep))]
    ws = openpyxl.load_workbook(chemin, read_only=True, data_only=True).worksheets[0]
    return [["" if c is None else str(c) for c in row[:MAX_COLONNES]]
            for _, row in zip(range(MAX_LIGNES), ws.iter_rows(values_only=True))]


def _grille(s, a: ArchiveRapport) -> tuple[list[list[str]], int]:
    if f".{a.format}" not in EDITABLES:
        raise HTTPException(422, f"Format .{a.format} non éditable en ligne (xlsx ou csv seulement).")
    grille = _grille_origine(a)
    modifs = s.execute(select(ArchiveDonnees).where(ArchiveDonnees.archive_id == a.id)
                       .order_by(ArchiveDonnees.id)).scalars().all()
    for m in modifs:                           # ordre chronologique : la dernière gagne
        col = int(m.colonne)
        while len(grille) <= m.ligne:
            grille.append([])
        ligne = grille[m.ligne]
        while len(ligne) <= col:
            ligne.append("")
        ligne[col] = m.valeur or ""
    largeur = max((len(l) for l in grille), default=0)
    return [l + [""] * (largeur - len(l)) for l in grille], len(modifs)


@routeur.get("/archives/{archive_id}/donnees")
def lire_donnees(archive_id: int, user: dict = Depends(utilisateur_courant)):
    exiger_role(user, ROLES_ACCES_TOTAL)
    s = get_session()
    try:
        a = _archive(s, archive_id)
        grille, n = _grille(s, a)
        derniere = s.execute(select(ArchiveDonnees.modifie_par, ArchiveDonnees.modifie_le)
                             .where(ArchiveDonnees.archive_id == archive_id)
                             .order_by(ArchiveDonnees.id.desc()).limit(1)).first()
        return {"archive": _vue(a), "grille": grille, "nb_modifications": n,
                "derniere_modification": ({"par": derniere[0], "le": derniere[1].isoformat()}
                                          if derniere else None)}
    finally:
        s.close()


class Modification(BaseModel):
    ligne: int
    colonne: int
    valeur: str


class Modifications(BaseModel):
    modifications: list[Modification]


@routeur.post("/archives/{archive_id}/donnees")
def modifier_donnees(archive_id: int, corps: Modifications, user: dict = Depends(utilisateur_courant)):
    exiger_role(user, ROLES_ECRITURE)
    if not corps.modifications:
        raise HTTPException(422, "Aucune modification envoyée.")
    for m in corps.modifications:
        if not (0 <= m.ligne < MAX_LIGNES and 0 <= m.colonne < MAX_COLONNES):
            raise HTTPException(422, f"Cellule hors limites ({m.ligne}, {m.colonne}) : "
                                     f"{MAX_LIGNES} lignes × {MAX_COLONNES} colonnes au plus.")
    s = get_session()
    try:
        a = _archive(s, archive_id)
        _grille(s, a)                                    # refuse les formats non éditables
        maintenant = dt.datetime.now()
        s.add_all(ArchiveDonnees(archive_id=archive_id, ligne=m.ligne, colonne=str(m.colonne),
                                 valeur=m.valeur[:2000], modifie_par=user["login"],
                                 modifie_le=maintenant) for m in corps.modifications)
        s.commit()
    finally:
        s.close()
    return {"enregistrees": len(corps.modifications)}


@routeur.get("/archives/{archive_id}/export")
def exporter(archive_id: int, user: dict = Depends(utilisateur_courant)):
    exiger_role(user, ROLES_ACCES_TOTAL)
    s = get_session()
    try:
        a = _archive(s, archive_id)
        grille, _ = _grille(s, a)
        titre, version = a.titre, a.version
    finally:
        s.close()
    wb = openpyxl.Workbook()
    ws = wb.active
    for ligne in grille:
        ws.append([_nombre(v) for v in ligne])
    tampon = io.BytesIO()
    wb.save(tampon)
    nom = "".join(c if c.isalnum() or c in "-_ " else "_" for c in titre)[:80]
    return Response(tampon.getvalue(), media_type=TYPE_XLSX, headers={
        "Content-Disposition": f'attachment; filename="{nom}_v{version}_modifie.xlsx"'})


def _nombre(v: str):
    """Un nombre saisi redevient un nombre dans l'export (sinon Excel ne sait pas le sommer)."""
    try:
        return float(v.replace(" ", "").replace("\xa0", "").replace(",", ".")) if v.strip() else None
    except ValueError:
        return v


# ─── Séries ──────────────────────────────────────────────────────────────────
@routeur.get("/series/indicateurs")
def indicateurs(user: dict = Depends(utilisateur_courant)):
    exiger_role(user, ROLES_ACCES_TOTAL)
    s = get_session()
    try:
        lignes = s.execute(select(SerieIndicateur.indicateur, SerieIndicateur.unite, func.count(),
                                  func.min(SerieIndicateur.date_arrete), func.max(SerieIndicateur.date_arrete))
                           .group_by(SerieIndicateur.indicateur, SerieIndicateur.unite)).all()
        agences = sorted(a for a in s.execute(select(SerieIndicateur.agence).distinct()).scalars() if a)
    finally:
        s.close()
    return {"indicateurs": [{"indicateur": i, "unite": u, "points": n, "du": d.isoformat(),
                             "au": f.isoformat()} for i, u, n, d, f in sorted(lignes)],
            "agences": agences, "arretes_non_alimentes": arretes_non_alimentes()}


@routeur.get("/series")
def serie(indicateur: str, agence: str | None = None, user: dict = Depends(utilisateur_courant)):
    exiger_role(user, ROLES_ACCES_TOTAL)
    s = get_session()
    try:
        q = select(SerieIndicateur).where(SerieIndicateur.indicateur == indicateur)
        q = q.where(SerieIndicateur.agence.is_(None) if not agence else SerieIndicateur.agence == agence)
        points = s.execute(q.order_by(SerieIndicateur.date_arrete)).scalars().all()
    finally:
        s.close()
    return {"indicateur": indicateur, "agence": agence or None,
            "points": [{"date": p.date_arrete.isoformat(), "valeur": p.valeur, "unite": p.unite,
                        "source": p.source} for p in points]}


@routeur.post("/series/import")
def importer(fichier: UploadFile = File(...), user: dict = Depends(utilisateur_courant)):
    exiger_role(user, ROLES_ECRITURE)
    import shutil
    import tempfile
    nom = _nom_sain(fichier.filename, (".xlsx", ".xlsm"))
    dossier = tempfile.mkdtemp(prefix="series_")
    try:
        chemin = os.path.join(dossier, nom)
        _ecrire_sur_disque(fichier, chemin)
        try:
            return importer_series(chemin)
        except ValueError as e:
            raise HTTPException(400, str(e))
    finally:
        shutil.rmtree(dossier, ignore_errors=True)


@routeur.post("/series/alimenter")
def alimenter(arrete: str = Query(...), user: dict = Depends(utilisateur_courant)):
    exiger_role(user, ROLES_ECRITURE)
    try:
        d = dt.date.fromisoformat(arrete)
    except ValueError:
        raise HTTPException(422, f"Date invalide : {arrete}")
    return alimenter_depuis_moteurs(d)
