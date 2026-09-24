"""
Primes des catégories hors « AC et SUP » — orchestration (chantier 4).

Les RÈGLES de calcul sont celles de engine/moteur_primes.py (validées au centime contre
CALCUL_PRIMES) : ce module ne les réécrit pas, il leur fournit leurs bases depuis les
sources de la plateforme et applique les taux confirmés par le CDG.

  1. Direction d'agence   : % du RÉSULTAT COMPTABLE de l'agence (compte_resultat_agence).
                            Chef d'agence 1 %, adjoint 0,5 %. Perte → 0.
  2. Direction générale   : % du résultat TOTAL (colonne MICROPOP du même fichier).
                            DG 1 %, DGA 0,6 %, DAF 0,3 %, Responsable régional 1 %. Perte → 0.
  3. Fonctions support    : même montant pour chaque agent support d'une agence, selon les
                            réalisations de L'AGENCE (pas des individus) :
                              + 5 $ si objectif de décaissement atteint (≥ 100 %)
                              + 10 $ si épargne ≥ 60 % de l'encours crédit de l'agence
                              + PAR : ≤ 3 % → 30 $ ; 3-5 % → 20 $ ; 5-7 % → 10 $ ; > 7 % → 0
                            Total agence = prime unitaire × effectif support (saisi à la main).
  4. Recouvrement         : agents 1 % (91-180 j) / 3 % (181-360 j) / 5 % (radié) de leur
                            montant recouvré ; responsable 0,3 % / 0,5 % / 1 % du TOTAL.

STRICTEMENT ADDITIF : aucun moteur existant n'est modifié.
"""
from __future__ import annotations

import openpyxl

from engine.moteur_primes import (prime_direction, prime_direction_agence,
                                  prime_fonction_support, prime_recouvrement)

# Taux confirmés par le CDG (septembre 2026). Versionnés par la campagne_prime qui les
# fige au moment du calcul (traçabilité) : un changement de taux ne réécrit pas le passé.
TAUX_DIRECTION_AGENCE = {"chef_agence": 0.01, "adjoint": 0.005}
TAUX_DIRECTION_GENERALE = {
    "Directeur Général": 0.01,
    "Directeur Général Adjoint": 0.006,
    "Directeur Administratif et Financier": 0.003,
    "Responsable régional": 0.01,
}


def _f(v) -> float:
    if v in (None, ""):
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    return float(str(v).replace("\xa0", "").replace(" ", "").replace(" ", "")
                 .replace(",", "."))


# ─────────────────────────────────────────────────────────────────────────────
# 1-2. DIRECTION
# ─────────────────────────────────────────────────────────────────────────────
def primes_direction_agences(resultats: dict[str, float]) -> list[dict]:
    """resultats = {agence: résultat comptable}. Une ligne par agence."""
    lignes = []
    for agence, res in sorted(resultats.items()):
        p = prime_direction_agence(res, TAUX_DIRECTION_AGENCE["chef_agence"],
                                   TAUX_DIRECTION_AGENCE["adjoint"])
        lignes.append({"agence": agence, "resultat": res,
                       "prime_chef_agence": p["prime_directeur"],
                       "prime_adjoint": p["prime_adjoint"],
                       "motif": "Résultat positif" if res > 0 else "Perte : pas de prime"})
    return lignes


def primes_direction_generale(resultat_total: float) -> dict:
    return {"resultat_total": resultat_total,
            "primes": prime_direction(resultat_total, TAUX_DIRECTION_GENERALE),
            "motif": "Résultat positif" if resultat_total > 0 else "Perte : pas de prime"}


# ─────────────────────────────────────────────────────────────────────────────
# 3. FONCTIONS SUPPORT (par agence)
# ─────────────────────────────────────────────────────────────────────────────
def primes_support(agences: list[dict]) -> dict:
    """agences = [{agence, taux_decaissement, par30, epargne, encours, effectif}].

    taux_decaissement : réalisé / objectif de l'agence (None = objectif inconnu → critère
    non atteint, et signalé) ; par30 en fraction (0,0334 = 3,34 %).
    """
    lignes = []
    for a in agences:
        couverture = (a["epargne"] / a["encours"]) if a.get("encours") else 0.0
        p = prime_fonction_support(a.get("taux_decaissement") or 0.0, a["par30"], couverture,
                                   a.get("effectif") or 0)
        lignes.append({
            "agence": a["agence"], "taux_decaissement": a.get("taux_decaissement"),
            "par30": a["par30"], "couverture": couverture, "effectif": a.get("effectif"),
            "prime_decaissement": p["prime_volume"], "prime_epargne": p["prime_couverture"],
            "prime_par": p["prime_par"], "prime_unitaire": p["prime_unitaire"],
            "prime_totale_agence": p["prime_totale"],
            "objectif_connu": a.get("taux_decaissement") is not None,
            "effectif_saisi": a.get("effectif") is not None,
        })
    return {"agences": lignes, "total": round(sum(l["prime_totale_agence"] for l in lignes), 2)}


# ─────────────────────────────────────────────────────────────────────────────
# 4. RECOUVREMENT
# ─────────────────────────────────────────────────────────────────────────────
ENTETES_RECOUVREMENT = ("equipe", "agent", "agence", "91-180", "181", "radie")


def _norm(t) -> str:
    import unicodedata
    t = unicodedata.normalize("NFKD", str(t or "")).encode("ascii", "ignore").decode()
    return t.strip().lower()


def lire_fichier_recouvrement(chemin: str, feuille: str | None = None) -> dict:
    """Tableau « Équipe | Agent | Agence | Montant 91-180 | Montant 181+ | Montant Radié ».

    La ligne TOTAL du fichier n'est pas une donnée : elle sert de CONTRÔLE (la somme des
    agents doit la redonner, sinon alerte — une ligne oubliée ou un total mal tiré).
    """
    wb = openpyxl.load_workbook(chemin, read_only=True, data_only=True)
    ws = wb[feuille] if feuille and feuille in wb.sheetnames else wb.worksheets[0]
    lignes, total_fichier, cols = [], None, None
    for row in ws.iter_rows(values_only=True):
        cellules = [_norm(c) for c in row]
        if cols is None:
            if any("agent" == c for c in cellules) and any("radi" in c for c in cellules):
                cols = {
                    "equipe": next(i for i, c in enumerate(cellules) if c.startswith("equipe")),
                    "agent": cellules.index("agent"),
                    "agence": next(i for i, c in enumerate(cellules) if c.startswith("agence")),
                    "m91": next(i for i, c in enumerate(cellules) if "91" in c),
                    "m181": next(i for i, c in enumerate(cellules) if "181" in c),
                    "radie": next(i for i, c in enumerate(cellules) if "radi" in c),
                }
            continue
        if not any(row):
            continue
        valeurs = (_f(row[cols["m91"]]), _f(row[cols["m181"]]), _f(row[cols["radie"]]))
        if any("total" in c for c in cellules if c):
            total_fichier = valeurs
            continue
        if not row[cols["agent"]]:
            continue
        lignes.append({"equipe": row[cols["equipe"]], "agent": str(row[cols["agent"]]).strip(),
                       "agence": row[cols["agence"]], "m91_180": valeurs[0],
                       "m181_360": valeurs[1], "radie": valeurs[2]})
    if cols is None:
        raise ValueError("En-tête introuvable : attendu Équipe | Agent | Agence | "
                         "Montant 91-180 | Montant 181+ | Montant Radié.")
    return {"lignes": lignes, "total_fichier": total_fichier}


def primes_recouvrement(lignes: list[dict], total_fichier=None) -> dict:
    agents = []
    for l in lignes:
        p = prime_recouvrement(l["m91_180"], l["m181_360"], l["radie"], agent=True)
        agents.append({**l, "prime_91_180": p["detail"]["91-180"],
                       "prime_181_360": p["detail"]["181+"],
                       "prime_radie": p["detail"]["radié"], "prime": p["prime"]})
    tot = (round(sum(l["m91_180"] for l in lignes), 2),
           round(sum(l["m181_360"] for l in lignes), 2),
           round(sum(l["radie"] for l in lignes), 2))
    resp = prime_recouvrement(*tot, agent=False)
    alertes = []
    if total_fichier is not None:
        for nom, calc, lu in zip(("91-180", "181+", "radié"), tot, total_fichier):
            if abs(calc - lu) > 0.01:
                alertes.append(f"Total {nom} : somme des agents {calc:,.2f} ≠ ligne TOTAL {lu:,.2f}")
    return {
        "agents": agents,
        "total_recouvre": {"91-180": tot[0], "181-360": tot[1], "radie": tot[2]},
        "total_primes_agents": round(sum(a["prime"] for a in agents), 2),
        "responsable": {"prime": resp["prime"], "detail": resp["detail"]},
        "alertes": alertes,
    }


# ─────────────────────────────────────────────────────────────────────────────
# BASES LUES DANS LA PLATEFORME (source unique : jamais ressaisies)
# ─────────────────────────────────────────────────────────────────────────────
def resultats_agences_depuis_base(date_arrete, db_path="socle/micropop.db") -> dict[str, float]:
    """RÉSULTAT COMPTABLE par agence, tel qu'importé (compte_resultat_agence)."""
    from sqlalchemy import select
    from socle.schema import CompteResultatAgence as CRA, get_session
    s = get_session(db_path)
    try:
        lignes = s.execute(select(CRA.poste, CRA.agence, CRA.montant)
                           .where(CRA.date_arrete == date_arrete)).all()
    finally:
        s.close()
    res = {ag: m for poste, ag, m in lignes
           if "RESULTAT" in poste.upper() and "COMPTABLE" in poste.upper()}
    if not res:
        raise ValueError(f"Aucun compte de résultat par agence importé pour {date_arrete} : "
                         "l'importer d'abord (page Import, domaine « compte de résultat par agence »).")
    return res


def bases_support(date_arrete, effectifs: dict[str, int], db_path="socle/micropop.db") -> dict:
    """Réalisations de chaque agence pour la prime support, depuis les moteurs existants :
    PAR30 et encours (engine.par), décaissements du mois (engine.decaissement), épargne
    convertie en USD au taux daté (même règle que engine.epargne), objectif de décaissement
    = Σ objectif_volume des agents de l'agence (param_objectif, date d'effet ≤ arrêté)."""
    from sqlalchemy import func, select
    from socle.schema import FaitEpargne, ParamObjectif, get_session
    from engine.par import calculer_par
    from engine.decaissement import decaissements
    from engine.etats_financiers import taux_change

    par = {a.designation: a for a in calculer_par(date_arrete, db_path=db_path)["agences"]}
    dec = decaissements(date_arrete, db_path=db_path)["par_agence"]
    s = get_session(db_path)
    try:
        taux = taux_change(s, date_arrete)
        ep_rows = s.execute(select(FaitEpargne.agence, FaitEpargne.devise,
                                   func.sum(FaitEpargne.solde_actuel))
                            .where(FaitEpargne.date_arrete == date_arrete)
                            .group_by(FaitEpargne.agence, FaitEpargne.devise)).all()
        effet = s.execute(select(func.max(ParamObjectif.date_effet))
                          .where(ParamObjectif.date_effet <= date_arrete)).scalar()
        obj_rows = [] if effet is None else s.execute(
            select(ParamObjectif.agence, func.sum(ParamObjectif.objectif_volume))
            .where(ParamObjectif.date_effet == effet).group_by(ParamObjectif.agence)).all()
    finally:
        s.close()
    if not ep_rows:
        raise ValueError(f"Aucune épargne importée pour {date_arrete} : critère épargne incalculable.")

    epargne: dict[str, float] = {}
    for ag, devise, solde in ep_rows:
        if ag:
            epargne[ag] = epargne.get(ag, 0.0) + ((solde or 0.0) / taux if devise == "CDF"
                                                  else (solde or 0.0))
    objectifs = {(ag or "").strip().upper(): v for ag, v in obj_rows}

    agences, alertes = [], []
    for nom, l in sorted(par.items()):
        obj = objectifs.get(nom.strip().upper())
        realise = dec.get(nom, {}).get("volume", 0.0)
        agences.append({"agence": nom, "par30": l.pct_par30, "encours": l.encours,
                        "epargne": epargne.get(nom, 0.0), "decaissement": realise,
                        "objectif_decaissement": obj,
                        "taux_decaissement": (realise / obj) if obj else None,
                        "effectif": effectifs.get(nom)})
        if not obj:
            alertes.append(f"{nom} : objectif de décaissement inconnu → critère 5 $ non accordé")
        if effectifs.get(nom) is None:
            alertes.append(f"{nom} : effectif support non saisi → total agence = 0")
    return {"agences": agences, "taux_change": taux, "date_effet_objectifs": effet,
            "alertes": alertes}
