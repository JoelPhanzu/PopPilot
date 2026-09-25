"""
Moteur : import du compte de résultat par agence (chantier 3).
Source : fichier mensuel du CDG (« COMPTE RESULTAT <MOIS> <ANNÉE> »), en .xlsx / .xlsm / .xls
(feuille Feuil2), .csv (séparateur ; , ou tabulation) ou .pdf (tableau texte, pas une image
scannée). Seul l'extracteur change : la lecture et le contrôle ci-dessous sont les mêmes.
Structure PILOTÉE PAR LES EN-TÊTES (et non par des positions fixes) :
  - colonne A (1) = intitulé du poste
  - ligne d'en-tête = celle qui porte « MICROPOP » ; les colonnes entre A et MICROPOP sont
    les agences, reconnues par leur nom (VICTOIRE, OZONE, GOMA, LUBUMBASHI, MASINA, GOMBE)
  - tout ce qui est À DROITE de MICROPOP est ignoré (certains fichiers y recopient le mois
    précédent)
  - lignes clés : "TOTAL PRODUITS", "TOTAL CHARGES", "RESULTAT COMPTABLE"
Le nombre d'agences varie dans le temps (4 en janvier 2025, Masina puis Gombe ensuite) :
une agence absente du fichier n'a simplement pas de ligne ce mois-là.
Une colonne d'agence INCONNUE → refus : on ne devine pas à quelle agence elle appartient.
CONTRÔLE : la colonne MICROPOP doit égaler la somme des agences (sinon alerte).
"""
from __future__ import annotations
import csv
import datetime as dt
import io
import os
import re

import openpyxl

# nom court de l'en-tête → nom de l'agence dans le socle
AGENCES = {"VICTOIRE": "AGENCE DE VICTOIRE", "OZONE": "AGENCE OZONE", "GOMA": "AGENCE DE GOMA",
           "LUBUMBASHI": "AGENCE DE LUBUMBASHI", "MASINA": "AGENCE DE MASINA",
           "GOMBE": "AGENCE DE GOMBE"}
CONSOLIDE = "MICROPOP"


def _f(v):
    if v in (None, ""):
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    t = re.sub(r"[\s\xa0\u202f$]|USD", "", str(v))
    negatif = t.startswith("(") and t.endswith(")")          # (1 234,56) = négatif (PDF)
    t = t.strip("()")
    if "," in t and "." in t:                                # 1.234,56 ou 1,234.56
        t = t.replace(".", "").replace(",", ".") if t.rfind(",") > t.rfind(".") else t.replace(",", "")
    else:
        t = t.replace(",", ".")
    try:
        x = float(t)
    except ValueError:
        return 0.0
    return -x if negatif else x


def lire_lignes(path, feuille="Feuil2") -> list[tuple]:
    """Toutes les lignes du fichier, quel que soit le format, en tuples de cellules."""
    ext = os.path.splitext(str(path))[1].lower()
    if ext in (".xlsx", ".xlsm"):
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[feuille] if feuille in wb.sheetnames else wb.worksheets[0]
        rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
        return rows
    if ext == ".xls":
        import xlrd
        wb = xlrd.open_workbook(path)
        ws = wb.sheet_by_name(feuille) if feuille in wb.sheet_names() else wb.sheet_by_index(0)
        return [tuple(ws.row_values(i)) for i in range(ws.nrows)]
    if ext == ".csv":
        brut = open(path, "rb").read()
        for enc in ("utf-8-sig", "cp1252"):
            try:
                texte = brut.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        dialecte = csv.Sniffer().sniff(texte[:4096], delimiters=";,\t")
        return [tuple(r) for r in csv.reader(io.StringIO(texte), dialecte)]
    if ext == ".pdf":
        import pdfplumber
        rows = []
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                for table in page.extract_tables():
                    rows.extend(tuple(c for c in r) for r in table)
        if not rows:
            raise ValueError("Aucun tableau lisible dans le PDF (document scanné ?). "
                             "Exporter le compte de résultat en PDF depuis Excel, ou en .xlsx / .csv.")
        return rows
    raise ValueError(f"Format non pris en charge : {ext} (attendu .xlsx, .xlsm, .xls, .csv, .pdf).")


def _nom(v) -> str:
    return " ".join(str(v or "").upper().split())


def _en_tete(rows) -> tuple[int, dict[int, str], int]:
    """(index de la ligne d'en-tête, {colonne 0-based: agence}, colonne MICROPOP 0-based)."""
    for i, row in enumerate(rows):
        noms = [_nom(v) for v in row]
        if CONSOLIDE not in noms[1:]:
            continue
        col_conso = noms.index(CONSOLIDE, 1)
        agences = {}
        for c in range(1, col_conso):
            n = noms[c].removeprefix("AGENCE DE ").removeprefix("AGENCE ").strip()
            if not n:
                continue
            if n not in AGENCES:
                raise ValueError(f"Colonne d'agence inconnue « {row[c]} » dans l'en-tête "
                                 f"(attendu : {', '.join(AGENCES)}).")
            agences[c] = AGENCES[n]
        if not agences:
            raise ValueError("En-tête sans aucune colonne d'agence avant MICROPOP.")
        return i, agences, col_conso
    raise ValueError("Ligne d'en-tête introuvable : aucune cellule « MICROPOP » (feuille du "
                     "compte de résultat par agence attendue).")


def importer_compte_resultat(path, date_arrete: dt.date, feuille="Feuil2",
                             tolerance=1.0):
    """Renvoie la liste des lignes (poste, agence, montant) + un rapport de contrôle."""
    rows = lire_lignes(path, feuille)
    i_entete, agences, col_conso = _en_tete(rows)

    lignes = []
    controles = []
    for row in rows[i_entete + 1:]:
        intitule = row[0] if row else None
        if not intitule or not str(intitule).strip():
            continue
        intitule = " ".join(str(intitule).split())   # espaces normalisés (PDF, CSV, Excel)
        if intitule.upper() in ("INTITULE",) or "SITUATION" in intitule.upper():
            continue  # en-têtes
        somme_agences = 0.0
        for col, agence in agences.items():
            montant = _f(row[col]) if len(row) > col else 0.0
            lignes.append({"date_arrete": date_arrete, "poste": intitule,
                           "agence": agence, "montant": montant})
            somme_agences += montant
        consolide = _f(row[col_conso]) if len(row) > col_conso else 0.0
        if abs(consolide - somme_agences) > tolerance:
            controles.append({"poste": intitule, "consolide": consolide,
                              "somme_agences": somme_agences,
                              "ecart": consolide - somme_agences})
    return {"lignes": lignes, "controles_incoherents": controles,
            "nb_lignes": len(lignes), "nb_alertes": len(controles),
            "agences": list(agences.values())}


def resultat_par_agence(lignes):
    """Extrait le RÉSULTAT COMPTABLE par agence depuis les lignes importées."""
    res = {}
    for l in lignes:
        if "RESULTAT" in l["poste"].upper() and "COMPTABLE" in l["poste"].upper():
            res[l["agence"]] = l["montant"]
    return res


if __name__ == "__main__":
    import sys
    r = importer_compte_resultat(sys.argv[1], dt.date(2026, 7, 31))
    print(f"Lignes importées : {r['nb_lignes']} | alertes de cohérence : {r['nb_alertes']}")
    res = resultat_par_agence(r["lignes"])
    print("\nRÉSULTAT COMPTABLE par agence :")
    for ag, m in sorted(res.items(), key=lambda x: -x[1]):
        print(f"  {ag:26s} : {m:>14,.2f}  {'bénéfice' if m >= 0 else 'PERTE'}")
