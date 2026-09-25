"""
Tests du tableau de bord crédit complet (engine/tableau_de_bord_credit.py).

Références : Dashboard DailyToolReporting de MAI 2026 (M-1 = avril), déjà validées moteur par
moteur dans test_phase1_par — ce tableau de bord doit les redonner AU CENTIME, puisqu'il ne fait
qu'assembler les mêmes moteurs sur un seul chargement :
  décaissements 1-31 mai 517 / 1 070 672 ; coût du risque 1 119,37 ; 120 entrées en PAR
  (116 482,40) ; migrations 5 tranches ; croissance −0,8904 % ; provisions 938 244,42 ;
  encours/PAR par agence (Dashboard).
Plus : Σ lignes = MICROPOP à chaque niveau ; P15 ; période de flux libre ; roster du mois.

Lancer :  python tests/test_tableau_de_bord_credit.py
"""
import datetime as dt
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import donnees_test as D   # noqa: E402

DB = "socle/test_tdb_credit.db"
MAI, AVRIL = dt.date(2026, 5, 30), dt.date(2026, 4, 30)
REF_AGENCES = {
    "AGENCE DE VICTOIRE": (2531178.47, 274877.29, 260281.47, 238210.18),
    "AGENCE OZONE": (1767802.20, 319713.96, 261424.23, 221564.02),
    "AGENCE DE GOMA": (1453955.88, 39926.86, 39402.33, 39402.33),
    "AGENCE DE LUBUMBASHI": (2027575.81, 240365.80, 230509.03, 208636.48),
    "AGENCE DE MASINA": (2027663.12, 222202.56, 181675.27, 158532.93),
    "AGENCE DE GOMBE": (1006155.18, 91360.75, 78825.72, 69563.83),
}
MIGR = {"31-60": 40814.15, "61-90": 69246.52, "91-180": 73868.45, "181-360": 55535.31,
        "361+": 66584.24}
_pret = []


def _tdb(**kw):
    if not _pret:
        import socle.schema as S
        os.environ.pop("DATABASE_URL", None)
        mai, avril = D.exiger("Enours_MAI_2026_.xls"), D.exiger("Encours_crédit_AVRIL_2026.xlsx")
        S.fermer_moteurs()
        if os.path.exists(DB):
            os.remove(DB)
        from socle.seed_parametres import seed
        from ingest.import_credit import importer_credit
        seed(DB)
        importer_credit(avril, AVRIL, date_snapshot=dt.date(2026, 5, 4), db_path=DB)
        importer_credit(mai, MAI, date_snapshot=dt.date(2026, 6, 1), db_path=DB)
        from socle.agences import seed_agences
        seed_agences(DB)                         # Goma FERMÉE, comme en production
        _pret.append(True)
    from engine.tableau_de_bord_credit import tableau_de_bord_credit
    kw.setdefault("debut", dt.date(2026, 5, 1))
    kw.setdefault("fin", dt.date(2026, 5, 31))
    return tableau_de_bord_credit(MAI, db_path=DB, **kw)


def test_micropop_egale_dashboard_mai():
    r = _tdb()
    g = r["lignes"][0]
    assert r["precedent"] == "2026-04-30"
    assert g["designation"] == "MICROPOP" and abs(g["encours"] - 10814330.66) < 0.01
    assert g["decaisse_nombre"] == 517 and abs(g["decaisse_volume"] - 1070672.00) < 0.5
    assert abs(g["cout_du_risque"] - 1119.37) < 0.02
    assert g["entree_par_nb"] == 120 and abs(g["entree_par_montant"] - 116482.40) < 0.02
    for tr, v in MIGR.items():
        assert abs(g["migration_vers"][tr] - v) < 0.05, tr
    assert abs(g["croissance"] - (-0.008904)) < 0.0001 and abs(g["encours_m1"] - 10911491.75) < 0.01
    assert abs(g["provisions"] - 938244.42) < 0.01
    assert abs(g["par30"] - 1052118.05) < 0.01
    # Variation de provision = provision à date − provision fin M-1 (= /provisions d'avril).
    from engine.derivation import deriver_provisions
    avril = deriver_provisions(AVRIL, db_path=DB)["provision_capital_totale"]
    assert abs(g["provisions_m1"] - avril) < 0.01, (g["provisions_m1"], avril)
    assert abs(g["variation_provision"] - (938244.42 - avril)) < 0.02


def test_agences_egalent_dashboard_et_somme():
    r = _tdb()
    g, agences = r["lignes"][0], {l["designation"]: l for l in r["lignes"][1:]}
    for nom, (enc, p1, p30, p90) in REF_AGENCES.items():
        l = agences[nom]
        for cle, v in (("encours", enc), ("par1", p1), ("par30", p30), ("par90", p90)):
            assert abs(l[cle] - v) < 0.01, (nom, cle, l[cle], v)
    for cle in ("encours", "decaisse_nombre", "decaisse_volume", "cout_du_risque",
                "entree_par_montant", "provisions", "provisions_m1", "variation_provision",
                "encours_m1", "nb_credits", "p15"):
        assert abs(sum(l[cle] for l in agences.values()) - g[cle]) < 0.05, cle


def test_periode_de_flux_libre_et_p15():
    q1 = _tdb(fin=dt.date(2026, 5, 15))["lignes"][0]
    q2 = _tdb(debut=dt.date(2026, 5, 16))["lignes"][0]
    tout = _tdb()
    assert (q1["decaisse_nombre"], q2["decaisse_nombre"]) == (134, 383)       # Dashboard
    assert q1["decaisse_nombre"] + q2["decaisse_nombre"] == tout["lignes"][0]["decaisse_nombre"]
    assert tout["lignes"][0]["p15"] == 134                                   # P15 = 1er→15
    assert abs(q1["encours"] - tout["lignes"][0]["encours"]) < 0.01          # le stock ne bouge pas
    serie = tout["decaissements_jour"]
    assert len(serie) == 31 and serie[-1]["cumul_nombre"] == 517


def test_niveaux_roster_et_filtres():
    r = _tdb(niveau="agent")
    assert r["message"] and len(r["lignes"]) == 1                            # pas de roster → pas de ligne
    from ingest.import_objectifs import importer_objectifs
    importer_objectifs(D.exiger("OBJECTIF.xlsx"), dt.date(2026, 5, 1), db_path=DB)
    for niveau in ("agence", "superviseur", "agent"):
        r = _tdb(niveau=niveau)
        g, reste = r["lignes"][0], r["lignes"][1:]
        assert abs(sum(l["encours"] for l in reste) - g["encours"]) < 0.05, niveau
        assert abs(sum(l["cout_du_risque"] for l in reste) - g["cout_du_risque"]) < 0.05, niveau
    agents = _tdb(niveau="agent")["lignes"]
    assert any(l["statut"] == "orphelin" for l in agents) and any(l["statut"] == "gele" for l in agents)
    g = _tdb()["lignes"][0]
    assert g["nb_agents"] == 39 and g["p15_objectif"] == int(g["objectif_nombre"] / 2 + 0.5)
    assert abs(g["productivite"] - g["decaisse_nombre"] / 39) < 1e-9
    f = _tdb(filtres={"sexe": "F"})
    assert f["objectifs_applicables"] is False and f["lignes"][0]["objectif_nombre"] is None
    h = _tdb(filtres={"sexe": "H"})
    assert f["lignes"][0]["decaisse_nombre"] + h["lignes"][0]["decaisse_nombre"] == 517


def test_niveau_client_et_comptage_des_clients():
    from socle.schema import get_session, FaitCredit
    r = _tdb(niveau="client", limite=100000)
    g, reste = r["lignes"][0], r["lignes"][1:]
    assert r["nb_lignes_total"] == len(reste) and len(reste) > 1000
    assert abs(sum(l["encours"] for l in reste) - g["encours"]) < 0.05
    assert abs(sum(l["cout_du_risque"] for l in reste) - g["cout_du_risque"]) < 0.05
    # Règle CDG : clients = noms distincts ; crédits = dossiers
    s = get_session(DB)
    noms = {" ".join((n or "").split()).upper() for (n,) in s.query(FaitCredit.nom_client)
            .filter(FaitCredit.date_arrete == MAI)}
    s.close()
    assert g["nb_clients"] == len(noms) and g["nb_credits"] == 7984
    court = _tdb(niveau="client", limite=10)
    assert len(court["lignes"]) == 11 and court["nb_lignes_total"] == r["nb_lignes_total"]
    encours = [l["encours"] for l in court["lignes"][1:]]
    assert encours == sorted(encours, reverse=True)                       # les plus gros d'abord


def test_potentiel_fin_de_mois():
    from types import SimpleNamespace as P
    from engine.potentiel import projeter
    arrete = dt.date(2026, 5, 5)
    prets = [
        P(numero_dossier="a", encours=100, jours_de_retard=10, date_deboursement=dt.date(2026, 1, 3),
          frequence="Mensuelle", date_fin_echeance=None),                  # retard : vieillit de 26 j
        P(numero_dossier="b", encours=100, jours_de_retard=0, date_deboursement=dt.date(2026, 4, 10),
          frequence="Mensuelle", date_fin_echeance=dt.date(2026, 12, 10)),  # échéance 10/05 → 21 j
        P(numero_dossier="c", encours=100, jours_de_retard=0, date_deboursement=dt.date(2026, 5, 1),
          frequence="Tous les 28 jours", date_fin_echeance=None),           # échéance 29/05 → 2 j
        P(numero_dossier="d", encours=100, jours_de_retard=0, date_deboursement=dt.date(2026, 3, 31),
          frequence="Mensuelle", date_fin_echeance=None),                   # échéance 31/05 : pas encore en retard
        P(numero_dossier="e", encours=100, jours_de_retard=0, date_deboursement=dt.date(2026, 4, 10),
          frequence="Mensuelle", date_fin_echeance=dt.date(2026, 5, 4)),    # soldé avant : rien
    ]
    j = {k: v.jours_de_retard for k, v in projeter(prets, arrete).items()}
    assert j == {"a": 36, "b": 21, "c": 2, "d": 0, "e": 0}, j
    # Sur l'extraction réelle (arrêté 30/05, un jour avant la fin du mois)
    r = _tdb(niveau="agence")
    g, reste = r["lignes"][0], r["lignes"][1:]
    assert g["potentiel_cout_du_risque"] >= g["cout_du_risque"] - 0.01     # le retard ne recule pas
    assert abs(sum(l["potentiel_cout_du_risque"] for l in reste) - g["potentiel_cout_du_risque"]) < 0.05
    assert sum(l["potentiel_migration_nb"] for l in reste) == g["potentiel_migration_nb"]


def test_roster_nom_court_et_agence_suspendue():
    from socle.roster import correspondances
    roster = {("KANDA RODDY", "AGENCE OZONE"), ("JEAN", "AGENCE X"), ("JEAN PAUL", "AGENCE X")}
    m = correspondances(roster, [("MBOLELA KANDA  RODDY", "AGENCE OZONE"), ("KANDA RODDY", "AGENCE X"),
                                 ("JEAN PAUL MUKENDI", "AGENCE X"), ("JEAN PAUL", "AGENCE X")])
    assert m[("MBOLELA KANDA RODDY", "AGENCE OZONE")] == ("KANDA RODDY", "AGENCE OZONE")
    assert ("KANDA RODDY", "AGENCE X") not in m                            # autre agence : non
    assert ("JEAN PAUL MUKENDI", "AGENCE X") not in m                      # deux candidats : on ne devine pas
    assert m[("JEAN PAUL", "AGENCE X")] == ("JEAN PAUL", "AGENCE X")       # l'identique prime
    # Goma SUSPENDUE (production) : portefeuille gelé, jamais orphelin ; encours inchangé
    from socle.schema import get_session
    from socle.agences import enregistrer_agence
    _tdb()
    s = get_session(DB)
    enregistrer_agence(s, "AGENCE DE GOMA", statut="SUSPENDUE")
    s.close()
    try:
        r = _tdb(niveau="superviseur")
        goma = [l for l in r["lignes"] if l["agence"] == "AGENCE DE GOMA"]
        assert [l["statut"] for l in goma] == ["gele"], goma
        assert abs(goma[0]["encours"] - REF_AGENCES["AGENCE DE GOMA"][0]) < 0.01
    finally:
        s = get_session(DB)
        enregistrer_agence(s, "AGENCE DE GOMA", statut="FERMEE")
        s.close()


def test_objectifs_format_sept_colonnes():
    import openpyxl
    import tempfile
    from ingest.import_objectifs import importer_objectifs
    from socle.schema import get_session, ParamObjectif, DimEmploye
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "OBJECTIF"
    ws.append(["AGENCE", "SUPERVISEUR", "AGENT DE CREDIT", "#NOMBRE A DECAISSE ", "VOLUME ",
               "PORTEFEUILLE ", "PAR"])
    ws.append(["AGENCE OZONE", "KANDA RODDY", "SABWA TSHIBANGU PATRICK", None, None, None, "5%"])
    ws.append(["AGENCE OZONE", "KANDA RODDY", "MWAMBA MAGLOIRE  Magloire ", 10, "20 000", None, 0.05])
    chemin = os.path.join(tempfile.mkdtemp(), "objectif.xlsx")
    wb.save(chemin)
    _tdb()
    effet = dt.date(2026, 9, 1)
    r = importer_objectifs(chemin, effet, db_path=DB)
    assert r == {"agents": 2, "superviseurs": 1, "objectifs": 2}, r
    s = get_session(DB)
    o = {x.agent: x for x in s.query(ParamObjectif).filter(ParamObjectif.date_effet == effet)}
    assert o["SABWA TSHIBANGU PATRICK"].objectif_decaissement_nombre == 0 and \
        abs(o["SABWA TSHIBANGU PATRICK"].objectif_par - 0.05) < 1e-12      # « 5% » → 0,05
    assert o["MWAMBA MAGLOIRE Magloire"].objectif_volume == 20000         # espaces normalisés
    s.query(ParamObjectif).filter(ParamObjectif.date_effet == effet).delete()
    s.query(DimEmploye).filter(DimEmploye.date_debut == effet).delete()
    s.commit()
    s.close()


def test_top_clients_et_arretes():
    import filtres_credit as F
    from socle.schema import get_session
    _tdb()
    origine = F.get_session
    F.get_session = lambda *a, **k: get_session(DB)
    try:
        u = {"login": "cdg", "role": "CDG", "agence": None}
        r = F.endpoint_clients_top(arrete=MAI.isoformat(), n=20, critere="encours", debut=None, fin=None,
                                   agence=None, sexe=None, produits=None, duree=None, agent=None,
                                   superviseur=None, user=u)
        assert len(r["meilleurs"]) == 20 and len(r["pires"]) == 20
        assert all(c["max_jours_retard"] == 0 for c in r["meilleurs"])     # meilleurs : aucun retard
        v = [c["valeur"] for c in r["meilleurs"]]
        assert v == sorted(v, reverse=True)
        p = [c["encours_retard"] for c in r["pires"]]
        assert p == sorted(p, reverse=True) and all(c["max_jours_retard"] > 0 for c in r["pires"])
        d = F.endpoint_clients_top(arrete=MAI.isoformat(), n=50, critere="decaissement",
                                   debut="2026-05-01", fin="2026-05-31", agence=None, sexe=None,
                                   produits=None, duree=None, agent=None, superviseur=None, user=u)
        assert d["periode"] == ["2026-05-01", "2026-05-31"] and len(d["meilleurs"]) == 50
        a = {x["date"] for x in F.endpoint_arretes_credit(user=u)["arretes"]}
        assert a == {"2026-04-30", "2026-05-30"}
        ag = F.endpoint_clients_top(arrete=MAI.isoformat(), n=10, critere="encours", debut=None, fin=None,
                                    agence=None, sexe=None, produits=None, duree=None, agent=None,
                                    superviseur=None, user={"login": "v", "role": "AGENCE",
                                                            "agence": "AGENCE DE VICTOIRE"})
        assert {c["agence"] for c in ag["meilleurs"] + ag["pires"]} == {"AGENCE DE VICTOIRE"}
    finally:
        F.get_session = origine


def test_exports_csv_xlsx():
    """L'export reprend l'écran : mêmes lignes, mêmes chiffres, CSV lisible par Excel (« ; »,
    virgule décimale, BOM) ; un rôle AGENCE n'exporte que sa ligne, sans champs réservés."""
    import io as _io
    from urllib.parse import urlencode
    import openpyxl
    from fastapi import HTTPException
    from starlette.requests import Request
    import export_tableaux as X
    import filtres_credit as F
    from socle.schema import get_session
    _tdb()
    cdg = {"login": "cdg", "role": "CDG", "agence": None}

    def appel(fonction, user=cdg, **params):
        fmt = params.pop("format", "csv")
        req = Request({"type": "http", "query_string": urlencode(params).encode(), "headers": []})
        return fonction(req, format=fmt, user=user)

    import engine.tableau_de_bord_credit as T
    origine, origine_t = F.get_session, T.get_session
    F.get_session = T.get_session = lambda *a, **k: get_session(DB)
    try:
        r = appel(X.export_tableau_credit, arrete=MAI.isoformat(), debut="2026-05-01", fin="2026-05-31")
        texte = r.body.decode("utf-8")
        assert texte.startswith("\ufeff") and "Encours;" in texte           # BOM : Excel lit l'UTF-8
        micropop = next(l for l in texte.splitlines() if l.startswith("MICROPOP;MICROPOP"))
        assert "10814330,66" in micropop and ";517;" in micropop          # = écran, au centime
        x = appel(X.export_tableau_credit, arrete=MAI.isoformat(), format="xlsx", niveau="agence")
        ws = openpyxl.load_workbook(_io.BytesIO(x.body)).active
        assert any(v and v[1] == "AGENCE DE VICTOIRE" for v in ws.iter_rows(values_only=True))
        t = appel(X.export_tableau_clients, arrete=MAI.isoformat(), n=30)
        assert t.body.decode("utf-8").count("meilleurs;") == 30
        p = appel(X.export_tableau_credit, arrete=MAI.isoformat(), format="pdf", niveau="agence")
        assert p.body[:5] == b"%PDF-" and p.media_type == "application/pdf"
        try:
            appel(X.export_tableau_credit, arrete=MAI.isoformat(), format="doc")
            raise AssertionError("format inconnu aurait dû être refusé (422)")
        except HTTPException as e:
            assert e.status_code == 422
        agence = {"login": "v", "role": "AGENCE", "agence": "AGENCE DE VICTOIRE"}
        a = appel(X.export_tableau_credit, user=agence, arrete=MAI.isoformat()).body.decode("utf-8")
        assert "AGENCE OZONE" not in a and "MICROPOP;" not in a
    finally:
        F.get_session, T.get_session = origine, origine_t


if __name__ == "__main__":
    code = D.lancer("Tableau de bord credit complet", [
        (test_exports_csv_xlsx, "Exports CSV / Excel / PDF = écran ; AGENCE limitée à sa ligne"),
        (test_micropop_egale_dashboard_mai, "MICROPOP mai = Dashboard (décaissements, CR, migrations, croissance, provisions)"),
        (test_agences_egalent_dashboard_et_somme, "Agences = Dashboard ; Σ agences = MICROPOP"),
        (test_periode_de_flux_libre_et_p15, "Flux 1-15 / 16-31 mai = 134 / 383 ; P15 ; stock inchangé"),
        (test_niveaux_roster_et_filtres, "Niveaux agent/superviseur (roster), objectifs, productivité, filtres"),
        (test_niveau_client_et_comptage_des_clients, "Niveau client : Σ = MICROPOP ; clients = noms distincts"),
        (test_potentiel_fin_de_mois, "Potentiel CR / migration : projection fin de mois"),
        (test_roster_nom_court_et_agence_suspendue, "Roster en nom court ; agence suspendue = gelée"),
        (test_objectifs_format_sept_colonnes, "Fichier OBJECTIF à 7 colonnes, cellules vides, « 5% »"),
        (test_top_clients_et_arretes, "Top N meilleurs / pires clients ; arrêtés disponibles"),
    ])
    from socle.schema import fermer_moteurs
    fermer_moteurs()
    D.sortir(code)
